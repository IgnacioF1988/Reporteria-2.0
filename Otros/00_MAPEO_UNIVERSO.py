# =============================================================================
# 00_MAPEO_UNIVERSO.py — Mapeo de universo multi-fondo (NUEVA LÓGICA MÉTRICAS)
# =============================================================================
# CONTEXTO
#   Reemplaza (en primera instancia) a INPUT_TD_BBG.xlsx como punto de partida.
#   INPUT_TD_BBG respondía "¿a qué proveedor voy a buscar la TABLA DE DESARROLLO?".
#   Este script responde una pregunta distinta y previa: "¿qué universo tengo,
#   en qué moneda debo salir (Hedge_Currency) y con qué atributos de familia
#   (REGS/144A/EMTN) cuento, para luego ir a buscar MÉTRICAS YA CALCULADAS
#   (Yield, Duration, ...) de BBG/RA por PK2?"
#
# INPUTS
#   CUBO_{CUBO_FECHA}.xlsx        (RUTA_CUBO_DIR)  → posiciones crudas
#   BD_INSTRUMENTOS.xlsx          (BD_INSTR_PATH)  → Risk_Currency, Investment_Type_Code (fuente canónica)
#   BD_FUNDS.xlsx                 (BD_FUNDS_PATH)  → FundShortName, FundBaseCurrency por ID_Fund
#
# OUTPUT
#   MAPEO_UNIVERSO_{CUBO_FECHA}.xlsx
#     - con_ISIN     : PK2 con ISIN (candidatos a BBG/RA)
#     - sin_ISIN     : PK2 sin ISIN (loans, distress, estructuras) + Sugerencia_Proveedor
#     - resumen      : conteos por fondo / proveedor sugerido / hedge
#
# REGLA Hedge_Currency (SOLO fondos HEDGE_FUNDS = [11, 17, 20])
#   Si Risk_Currency ∈ STRONG_CCY (USD, GBP, EUR) → se asume el papel está
#   hedgeado a la moneda local del Risk_Country (regla rápida, sujeta a
#   OVERRIDE manual posterior). Mapeo Risk_Country → moneda local en
#   RISK_COUNTRY_TO_LOCAL_CCY.
#   Para el resto de los fondos: Hedge_Currency = lo que traiga el CUBO (o NaN).
#
# NOTA IMPORTANTE
#   Filtra Investment_Type_Code == 1 (solo Renta Fija), igual que 00_CARTERA_VF.
#   Si se requiere el universo SIN ese filtro, comentar el bloque [FILTRO RF].
# =============================================================================

import os
import sys
import glob
import re
import numpy as np
import pandas as pd
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

_script_dir = os.path.dirname(os.path.abspath(__file__)) if '__file__' in dir() else os.getcwd()
if _script_dir not in sys.path: sys.path.insert(0, _script_dir)
if os.getcwd() not in sys.path:  sys.path.insert(0, os.getcwd())

from pipeline_config import (
    RUTA_TRABAJO, RUTA_CUBO_DIR, BD_INSTR_PATH, BD_FUNDS_PATH,
)

# ============================================================
# ⚙️ CONFIGURACIÓN — editable
# ============================================================
CUBO_FECHA   = "20260731"          # <<< fecha específica del CUBO a usar
FONDOS       = [65, 13, 59, 16, 17, 20, 68, 11, 2]   # <<< universo de fondos

HEDGE_FUNDS  = {11, 17, 20}        # fondos con lógica SWAP/NDF USD(u otra)->Local
STRONG_CCY   = {'USD', 'GBP', 'EUR'}

RISK_COUNTRY_TO_LOCAL_CCY = {
    'AR': 'ARS', 'BR': 'BRL', 'CL': 'CLP', 'PE': 'PEN',
    'UY': 'UYU', 'CO': 'COP', 'MX': 'MXN',
}

# Países donde SÍ tenemos RiskAmérica como proveedor alternativo (sin ISIN)
PAISES_RA = {'CL'}

# Columna de valor: rango dinámico [LocalPrice ... TotalMVal] del propio CUBO
COL_VALOR_INICIO = 'LocalPrice'
COL_VALOR_FIN    = 'TotalMVal'
COLS_VALOR_EXCLUIR = {'TotalMVal_Balance'}   # redundante

APLICAR_FILTRO_RF = True   # [FILTRO RF] Investment_Type_Code == 1

FNAME_OUT = os.path.join(RUTA_TRABAJO, f"MAPEO_UNIVERSO_{CUBO_FECHA}.xlsx")
# ============================================================

# ── Regex sufijos de serie (idéntico a 00_CARTERA_VF.py) ─────────────────────
SERIES_RE = re.compile(r'\s+(REGS|REGS-S|REG\s*S|144A|144a|144@|EMTN)\s*$', re.IGNORECASE)

def extract_base(name):
    return SERIES_RE.sub('', str(name).strip()).strip()

# ── Helpers formato Excel ─────────────────────────────────────────────────────
def fmt_ws(ws, hex_color, freeze='B2'):
    fill = PatternFill('solid', fgColor=hex_color)
    font = Font(color='FFFFFF', bold=True, size=9, name='Calibri')
    aln  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for cell in ws[1]:
        cell.fill = fill; cell.font = font; cell.alignment = aln
    if freeze:
        ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 28
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 8), 45)

print("=" * 70)
print(f"00_MAPEO_UNIVERSO | CUBO={CUBO_FECHA} | Fondos={FONDOS}")
print("=" * 70)

# =============================================================================
# [1] LEER CUBO ESPECÍFICO
# =============================================================================
print(f"\n[1] Buscando CUBO_{CUBO_FECHA}.xlsx...")

cubo_path = os.path.join(RUTA_CUBO_DIR, f"CUBO_{CUBO_FECHA}.xlsx")
if not os.path.exists(cubo_path):
    # fallback: buscar variantes con glob si el nombre exacto no calza
    candidatos = sorted(glob.glob(os.path.join(RUTA_CUBO_DIR, f"CUBO_{CUBO_FECHA}*.xlsx")))
    if not candidatos:
        print(f"   ERROR: no se encontró CUBO para {CUBO_FECHA} en:\n   {RUTA_CUBO_DIR}")
        sys.exit(1)
    cubo_path = candidatos[-1]

print(f"   Usando: {os.path.basename(cubo_path)}")
df_cubo = pd.read_excel(cubo_path, engine='openpyxl')
print(f"   Filas CUBO totales: {len(df_cubo):,}")
print(f"   Columnas CUBO ({len(df_cubo.columns)}): {list(df_cubo.columns)}")

# ID_Instrumento derivado del PK2 (parte antes del guión) — igual que 00_CARTERA
df_cubo['ID_Instrumento'] = pd.to_numeric(
    df_cubo['PK2'].astype(str).str.split('-').str[0], errors='coerce')

# Filtro de fondos
df_cubo = df_cubo[df_cubo['ID_Fund'].isin(FONDOS)].copy()
print(f"   Filas tras filtro fondos {FONDOS}: {len(df_cubo):,}")
faltan_fondos = sorted(set(FONDOS) - set(df_cubo['ID_Fund'].unique()))
if faltan_fondos:
    print(f"   [WARN] Fondos sin filas en el CUBO: {faltan_fondos}")

# =============================================================================
# [2] CRUZAR CON BD_INSTRUMENTOS (fuente canónica Risk_Currency / Investment_Type_Code)
# =============================================================================
print(f"\n[2] Cruzando con BD_INSTRUMENTOS...")

bd_instr = pd.read_excel(BD_INSTR_PATH, sheet_name='BD_INSTRUMENTOS', engine='openpyxl')
bd_instr['PK2'] = (bd_instr['ID_Instrumento'].astype(str).str.strip()
                   + '-' + bd_instr['SubID_Instrumento'].astype(str).str.strip())

bd_cols = ['PK2', 'Risk_Currency', 'Investment_Type_Code', 'Issue_Type_Code',
           'Name_Instrumento', 'ISIN', 'Risk_Country', 'CompanyName']
bd_cols = [c for c in bd_cols if c in bd_instr.columns]
faltantes_bd = [c for c in ('Name_Instrumento', 'ISIN', 'Risk_Country', 'CompanyName', 'Issue_Type_Code') if c not in bd_cols]
if faltantes_bd:
    print(f"   [WARN] Columnas no encontradas en BD_INSTRUMENTOS: {faltantes_bd}. Disponibles: {list(bd_instr.columns)}")

for _drop in ['Risk_Currency', 'Investment_Type_Code', 'Issue_Type_Code', 'Name_Instrumento', 'ISIN', 'Risk_Country', 'CompanyName']:
    if _drop in df_cubo.columns:
        df_cubo = df_cubo.drop(columns=[_drop])

df_cubo = pd.merge(df_cubo, bd_instr[bd_cols], on='PK2', how='left')
n_sin_rc = df_cubo['Risk_Currency'].isna().sum()
print(f"   PK2 sin Risk_Currency tras cruce: {n_sin_rc}")

# [FILTRO RF]
if APLICAR_FILTRO_RF and 'Investment_Type_Code' in df_cubo.columns:
    n_antes = len(df_cubo)
    df_cubo = df_cubo[df_cubo['Investment_Type_Code'] == 1].copy()
    print(f"   Filtro Investment_Type_Code=1: {n_antes:,} → {len(df_cubo):,}")

# Un registro por PK2 × ID_Fund
df_cubo = df_cubo.drop_duplicates(subset=['ID_Fund', 'PK2']).copy()
print(f"   PK2 únicos (fondo × PK2): {len(df_cubo):,}")

# =============================================================================
# [3] CRUZAR CON BD_FUNDS (FundShortName, FundBaseCurrency)
# =============================================================================
print(f"\n[3] Cruzando con BD_FUNDS...")

bd_funds = pd.read_excel(BD_FUNDS_PATH, engine='openpyxl')

# Detección tolerante de nombre de columna 'FundShortName'
col_short = next((c for c in bd_funds.columns if 'shortname' in c.lower().replace('_', '')), None)
col_basec = next((c for c in bd_funds.columns if 'basecurrency' in c.lower().replace('_', '')), None)

if col_short is None or col_basec is None:
    print(f"   [WARN] No se encontraron columnas esperadas en BD_FUNDS. Columnas disponibles: {list(bd_funds.columns)}")
if col_short is None: col_short = 'FundShortName'
if col_basec is None: col_basec = 'FundBaseCurrency'

bd_funds_cols = ['ID_Fund'] + [c for c in {col_short, col_basec} if c in bd_funds.columns]
bd_funds_slim = bd_funds[bd_funds_cols].drop_duplicates(subset=['ID_Fund']).rename(
    columns={col_short: 'FundShortName', col_basec: 'FundBaseCurrency'}
)

df_cubo = pd.merge(df_cubo, bd_funds_slim, on='ID_Fund', how='left')
n_sin_fund = df_cubo['FundShortName'].isna().sum() if 'FundShortName' in df_cubo.columns else len(df_cubo)
print(f"   PK2 sin match en BD_FUNDS: {n_sin_fund}")

# =============================================================================
# [4] LÓGICA DE FAMILIA — base_name / serie_hermanos (idéntica a 00_CARTERA_VF)
# =============================================================================
print(f"\n[4] Agrupamiento por familia (REGS/144A/EMTN)...")

df_cubo['Name_Instrumento'] = df_cubo['Name_Instrumento'].astype(str).str.strip()
df_cubo['base_name']    = df_cubo['Name_Instrumento'].apply(extract_base)
df_cubo['tiene_sufijo'] = df_cubo['Name_Instrumento'] != df_cubo['base_name']

# Grupo de familia a nivel de ISIN único (independiente del fondo que lo tenga)
isin_ok = df_cubo['ISIN'].notna() & (df_cubo['ISIN'].astype(str).str.strip() != '')
dedup_isin = df_cubo[isin_ok].drop_duplicates(subset=['ISIN'])[['ISIN', 'Name_Instrumento', 'base_name']]

grupos = (dedup_isin.groupby('base_name')
          .agg(n_series=('ISIN', 'count'), isins=('ISIN', list))
          .reset_index())
grupos_multi = grupos[grupos['n_series'] > 1]

serie_grupo_id_map = {}
serie_hermanos_map = {}
for _, g in grupos_multi.iterrows():
    isins = g['isins']
    gid   = isins[0]
    for isin in isins:
        serie_grupo_id_map[isin] = gid
        serie_hermanos_map[isin] = ','.join([x for x in isins if x != isin])

df_cubo['serie_grupo_id'] = df_cubo['ISIN'].map(serie_grupo_id_map).fillna('')
df_cubo['serie_hermanos'] = df_cubo['ISIN'].map(serie_hermanos_map).fillna('')

print(f"   Grupos multi-serie (ISIN único): {len(grupos_multi)}")

# =============================================================================
# [5] Hedge_Currency — regla rápida SOLO para HEDGE_FUNDS
# =============================================================================
print(f"\n[5] Aplicando regla Hedge_Currency para fondos {sorted(HEDGE_FUNDS)}...")

if 'Hedge_Currency' not in df_cubo.columns:
    df_cubo['Hedge_Currency'] = np.nan

def _regla_hedge(row):
    if row['ID_Fund'] not in HEDGE_FUNDS:
        return row.get('Hedge_Currency', np.nan)
    rc = str(row.get('Risk_Currency', '')).strip().upper()
    if rc in STRONG_CCY:
        pais = str(row.get('Risk_Country', '')).strip().upper()
        return RISK_COUNTRY_TO_LOCAL_CCY.get(pais, np.nan)
    return np.nan

df_cubo['Hedge_Currency_regla'] = df_cubo.apply(_regla_hedge, axis=1)

mask_hedge_funds = df_cubo['ID_Fund'].isin(HEDGE_FUNDS)
n_hedge_detectados = df_cubo.loc[mask_hedge_funds, 'Hedge_Currency_regla'].notna().sum()
print(f"   PK2 en HEDGE_FUNDS con Hedge_Currency derivado por regla: {n_hedge_detectados}")

pais_sin_map = (df_cubo.loc[mask_hedge_funds & df_cubo['Risk_Currency'].astype(str).str.upper().isin(STRONG_CCY)
                            & df_cubo['Hedge_Currency_regla'].isna(), 'Risk_Country']
                .dropna().unique().tolist())
if pais_sin_map:
    print(f"   [WARN] Risk_Country en strong-ccy sin mapeo local: {pais_sin_map} — completar RISK_COUNTRY_TO_LOCAL_CCY")

# Hedge_Currency final: para HEDGE_FUNDS, la regla manda; para el resto, se mantiene lo que trajo el CUBO
df_cubo['Hedge_Currency_final'] = np.where(
    mask_hedge_funds, df_cubo['Hedge_Currency_regla'], df_cubo['Hedge_Currency']
)

# =============================================================================
# [6] Columnas de valor dinámicas [LocalPrice ... TotalMVal]
# =============================================================================
print(f"\n[6] Detectando columnas de valor [{COL_VALOR_INICIO} .. {COL_VALOR_FIN}]...")

cols_all = list(df_cubo.columns)
cols_valor = []
if COL_VALOR_INICIO in cols_all and COL_VALOR_FIN in cols_all:
    i0 = cols_all.index(COL_VALOR_INICIO)
    i1 = cols_all.index(COL_VALOR_FIN)
    if i1 >= i0:
        cols_valor = [c for c in cols_all[i0:i1+1] if c not in COLS_VALOR_EXCLUIR]
    else:
        print(f"   [WARN] Orden inesperado: {COL_VALOR_FIN} aparece antes de {COL_VALOR_INICIO} en el CUBO")
else:
    faltan = [c for c in (COL_VALOR_INICIO, COL_VALOR_FIN) if c not in cols_all]
    print(f"   [WARN] Columnas no encontradas en el CUBO: {faltan}. Revisar nombres reales.")

print(f"   Columnas de valor detectadas: {cols_valor}")

# =============================================================================
# [7] ARMAR OUTPUT FINAL
# =============================================================================
print(f"\n[7] Armando estructura final...")

df_cubo['Hedge_Currency'] = df_cubo['Hedge_Currency_final']

COLS_FRONT = [
    'ID_Fund', 'FundShortName', 'FundBaseCurrency',
    'PK2', 'ID_Instrumento', 'ISIN', 'Name_Instrumento', 'CompanyName',
    'base_name', 'serie_grupo_id', 'serie_hermanos', 'tiene_sufijo',
    'Risk_Country', 'Risk_Currency', 'Hedge_Currency',
    'Investment_Type_Code', 'Issue_Type_Code',
]
COLS_FRONT = [c for c in COLS_FRONT if c in df_cubo.columns]
cols_finales = COLS_FRONT + [c for c in cols_valor if c not in COLS_FRONT]
cols_finales = [c for c in cols_finales if c in df_cubo.columns]

df_final = df_cubo[cols_finales].copy()

mask_isin = df_final['ISIN'].notna() & (df_final['ISIN'].astype(str).str.strip() != '')
df_con_isin = df_final[mask_isin].copy()
df_sin_isin = df_final[~mask_isin].copy()

# Sugerencia de proveedor para sin_ISIN
# RA solo aplica a bonos (Issue_Type_Code == 3, dentro del universo ya
# filtrado a Investment_Type_Code == 1). Esto acota el universo que se manda
# a revisar en el plug-in de RA — no todo Risk_Country=='CL' es bono.
def _sugerir_proveedor(row):
    pais = str(row.get('Risk_Country', '')).strip().upper()
    ityp = row.get('Issue_Type_Code', np.nan)
    es_bono = pd.notna(ityp) and int(ityp) == 3
    if pais in PAISES_RA and es_bono:
        return 'RA (posible) / revisar'
    return 'PROP'

df_sin_isin['Sugerencia_Proveedor'] = df_sin_isin.apply(_sugerir_proveedor, axis=1)

print(f"   con_ISIN : {len(df_con_isin):,} PK2")
print(f"   sin_ISIN : {len(df_sin_isin):,} PK2")

# =============================================================================
# [8] RESUMEN
# =============================================================================
print(f"\n[8] Generando resumen...")

resumen_rows = []
resumen_rows.append({'Métrica': '── UNIVERSO ──', 'Valor': '', 'Detalle': ''})
resumen_rows.append({'Métrica': 'PK2 total (fondo × PK2)', 'Valor': len(df_final), 'Detalle': ''})
resumen_rows.append({'Métrica': 'PK2 con ISIN', 'Valor': len(df_con_isin), 'Detalle': ''})
resumen_rows.append({'Métrica': 'PK2 sin ISIN', 'Valor': len(df_sin_isin), 'Detalle': ''})
resumen_rows.append({'Métrica': '', 'Valor': '', 'Detalle': ''})
resumen_rows.append({'Métrica': '── POR FONDO ──', 'Valor': '', 'Detalle': ''})
for fid in FONDOS:
    sub = df_final[df_final['ID_Fund'] == fid]
    nombre = sub['FundShortName'].iloc[0] if len(sub) and 'FundShortName' in sub.columns and sub['FundShortName'].notna().any() else '?'
    resumen_rows.append({'Métrica': f'Fondo {fid} ({nombre})', 'Valor': len(sub), 'Detalle': f"con_ISIN={sub['ISIN'].notna().sum() if len(sub) else 0}"})
resumen_rows.append({'Métrica': '', 'Valor': '', 'Detalle': ''})
resumen_rows.append({'Métrica': '── SIN ISIN — SUGERENCIA PROVEEDOR ──', 'Valor': '', 'Detalle': ''})
if not df_sin_isin.empty:
    for sug, n in df_sin_isin['Sugerencia_Proveedor'].value_counts().items():
        resumen_rows.append({'Métrica': sug, 'Valor': n, 'Detalle': ''})
resumen_rows.append({'Métrica': '', 'Valor': '', 'Detalle': ''})
resumen_rows.append({'Métrica': '── HEDGE_CURRENCY (fondos HEDGE_FUNDS) ──', 'Valor': '', 'Detalle': ''})
for fid in sorted(HEDGE_FUNDS):
    sub = df_final[df_final['ID_Fund'] == fid]
    n_hedge = sub['Hedge_Currency'].notna().sum() if len(sub) else 0
    resumen_rows.append({'Métrica': f'Fondo {fid} — con Hedge_Currency derivado', 'Valor': n_hedge, 'Detalle': f"total={len(sub)}"})
resumen_rows.append({'Métrica': '', 'Valor': '', 'Detalle': ''})
resumen_rows.append({'Métrica': '── FAMILIA (REGS/144A/EMTN) ──', 'Valor': '', 'Detalle': ''})
resumen_rows.append({'Métrica': 'Grupos multi-serie detectados', 'Valor': len(grupos_multi), 'Detalle': ''})
resumen_rows.append({'Métrica': 'ISIN con hermano de serie', 'Valor': df_final['serie_hermanos'].astype(str).str.strip().ne('').sum(), 'Detalle': ''})

df_resumen = pd.DataFrame(resumen_rows)

# =============================================================================
# [9] GUARDAR EXCEL
# =============================================================================
print(f"\n[9] Guardando: {FNAME_OUT}")

C_ISIN     = '0B2447'  # azul oscuro
C_SIN_ISIN = 'B8540B'  # naranja
C_RESUMEN  = '2E4057'  # gris azulado

with pd.ExcelWriter(FNAME_OUT, engine='openpyxl') as writer:
    df_con_isin.to_excel(writer, sheet_name='con_ISIN', index=False)
    fmt_ws(writer.sheets['con_ISIN'], C_ISIN)

    df_sin_isin.to_excel(writer, sheet_name='sin_ISIN', index=False)
    fmt_ws(writer.sheets['sin_ISIN'], C_SIN_ISIN)

    df_resumen.to_excel(writer, sheet_name='resumen', index=False)
    fmt_ws(writer.sheets['resumen'], C_RESUMEN, freeze='A2')

print("\n" + "=" * 70)
print("00_MAPEO_UNIVERSO completado.")
print(f"  con_ISIN : {len(df_con_isin):,}")
print(f"  sin_ISIN : {len(df_sin_isin):,}")
print(f"  Output   : {FNAME_OUT}")
print("=" * 70)