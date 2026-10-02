"""Excel final del cierre. Yields en decimal con formato de porcentaje solo en la celda."""
from __future__ import annotations

from pathlib import Path

import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
COLS_CARTERA = ["Pos_ID", "Pos_Key", "ID_Fund", "Fondo", "PK2", "PK2_Maestro", "BalanceSheet", "Name_Instrumento", "ISIN", "Risk_Country",
                "Risk_Currency", "Moneda_PK2", "Investment_Type_Code", "Issue_Type_Code", "Coupon_Type_Code",
                "Base_Name", "Familia", "ISIN_Hermanos", "Hedge_Currency", "Hedge_Origen", "Indice", "Overrides",
                "BalSheetKey", "Bucket", "Bucket_Origen", "Bucket_Orden", "Ficha_FI", "Ficha_Origen", "FX_Exposure", "FX_Origen", "Tratamiento",
                "Yield", "Duration", "Yield_Moneda", "Conversion", "Yield_Papel", "Duration_Papel", "Yield_XCCY", "Yield_Drop",
                "Dif_XCCY_Drop_bps", "Indice_Origen", "Extrapolado", "Fuente", "Origen", "Etapa", "Estado", "Estado_DEF", "CalcType",
                "CalcType_exportable", "Motivo", "Pedido_BBG", "TotalMVal", "MVBook", "AI", "LocalPrice", "Qty", "Factor"]


def _formatear(ws, pct_cols: set[str]):
    for c in ws[1]:
        c.fill, c.font = PatternFill("solid", fgColor=NAVY), Font(color="FFFFFF", bold=True, size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for j, col in enumerate(ws.iter_cols(1, ws.max_column), 1):
        ancho = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 9), 45)
        if col[0].value in pct_cols:
            for c in col[1:]:
                c.number_format = "0.00%"


def escribir_excel(hojas: dict[str, pd.DataFrame], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            (df if len(df) else pd.DataFrame({"info": ["sin filas"]})).to_excel(w, sheet_name=nombre[:31], index=False)
            _formatear(w.sheets[nombre[:31]], {"Yield", "Yield_Local", "Yield_Drop", "Yield_Papel", "Yield_XCCY", "AW", "DW", "Peso_MV", "Cobertura"})
    return path


def comparar_carteras(a: pd.DataFrame, b: pd.DataFrame, tol: float = 1e-9) -> pd.DataFrame:
    """Diferencias entre dos `cartera_final` por Pos_ID: solo en uno, o Yield/Duration/Fuente/Conversion/Estado/Hedge_Currency distintos."""
    a, b = a.drop_duplicates("Pos_ID").set_index("Pos_ID"), b.drop_duplicates("Pos_ID").set_index("Pos_ID")
    filas = []
    solo_a, solo_b = sorted(set(a.index) - set(b.index)), sorted(set(b.index) - set(a.index))
    if solo_a and solo_b and "Pos_Key" in a.columns and "Pos_Key" in b.columns:     # PK2 cambió de moneda: se cruza por Pos_Key
        kb = b.loc[solo_b].drop_duplicates("Pos_Key").reset_index().set_index("Pos_Key")["Pos_ID"]
        renombres = {pa: kb[a.at[pa, "Pos_Key"]] for pa in solo_a if a.at[pa, "Pos_Key"] in kb.index}
        for pa, pb in renombres.items():
            filas.append(dict(Pos_ID=pa, Campo="Pos_ID", A=pa, B=pb))
            b = b.rename(index={pb: pa})
        solo_a = [p for p in solo_a if p not in renombres]
        solo_b = [p for p in solo_b if p not in renombres.values()]
    for pid in solo_a + solo_b:
        filas.append(dict(Pos_ID=pid, Campo="Pos_ID", A="presente" if pid in a.index else "", B="presente" if pid in b.index else ""))
    comunes = a.index.intersection(b.index)
    for c in ("Yield", "Duration"):
        if c in a.columns and c in b.columns:
            x, y = pd.to_numeric(a.loc[comunes, c], errors="coerce"), pd.to_numeric(b.loc[comunes, c], errors="coerce")
            m = ~((x - y).abs() <= tol) & ~(x.isna() & y.isna())
            filas += [dict(Pos_ID=p, Campo=c, A=x[p], B=y[p]) for p in comunes[m.to_numpy()]]
    for c in ("Fuente", "Conversion", "Estado", "Hedge_Currency"):
        if c in a.columns and c in b.columns:
            x, y = a.loc[comunes, c].fillna("").astype(str), b.loc[comunes, c].fillna("").astype(str)
            filas += [dict(Pos_ID=p, Campo=c, A=x[p], B=y[p]) for p in comunes[(x != y).to_numpy()]]
    return pd.DataFrame(filas, columns=["Pos_ID", "Campo", "A", "B"])


# ── Plantilla de cajas: filas de REGLAS/cajas para las posiciones de caja sin regla, con sugerencias trazables ─────────
_NOMBRE_TASA = re.compile(r"^(?:CR|RD)_[A-Z]{3}_.+?_(\d{8})_(\d+(?:\.\d+)?)$")
SIN_INDICE_PLANTILLA = {"", "N.A.", "NA", "N/A", "NAN", "NONE"}


def _fecha_nombre(txt: str) -> pd.Timestamp | None:
    for fmt in ("%Y%m%d", "%d%m%Y"):                       # CR_CLP_SCOTIA_20260909_… y RD_USD_ERNCIV_20082026_…
        try:
            f = pd.Timestamp(pd.to_datetime(txt, format=fmt))
            if 2000 <= f.year <= 2100:
                return f
        except (ValueError, TypeError):
            continue
    return None


def plantilla_cajas(pos: pd.DataFrame, cajas: pd.DataFrame, settle: pd.Timestamp) -> pd.DataFrame:
    """Una fila por posición de caja que quedó sin fila en REGLAS/cajas, en el formato de esa hoja, con una sugerencia:
    COPIA_PK2 (mismo PK2 con fila en otro fondo), NOMBRE (tasa mensual y vencimiento en el nombre, CR_/RD_),
    POR_MONEDA (índice y spread más usados en las filas existentes de la misma familia de nombre —JPMCC, JPM, BNP…— y moneda)
    o vacío (cuenta corriente, fondo money market, simultánea: completar o dejar vacío = yield 0). Las columnas `_` son de referencia."""
    sin = pos[pos["Fuente"].eq("CAJA") & pos["Origen"].eq("SIN_REGLA")].copy()
    cols = ["ID_Fund", "PK2", "Indice_Referencia", "Spread_Anual", "Dias", "Comentario", "_Sugerencia", "_Fondo", "_Nombre", "_Bucket", "_Risk_Currency", "_TotalMVal"]
    if sin.empty:
        return pd.DataFrame(columns=cols)
    cj = cajas.copy() if cajas is not None and len(cajas) else pd.DataFrame(columns=["ID_Fund", "PK2", "Indice_Referencia", "Spread_Anual", "Dias"])
    cj["PK2"] = cj["PK2"].astype(str).str.strip()
    cj["_idx"] = cj["Indice_Referencia"].fillna("").astype(str).str.strip()      # NaN (pandas str dtype) no es un índice
    con_valor = cj[~cj["_idx"].str.upper().isin(SIN_INDICE_PLANTILLA) | cj["Spread_Anual"].notna()]
    por_pk2 = {r["PK2"]: r for _, r in con_valor.iterrows()}
    ref = pos.drop_duplicates("PK2").set_index("PK2")
    ccy_pos = ref["Risk_Currency"].astype(str).str.upper()
    fam_pos = ref["Name_Instrumento"].fillna("").astype(str).str.upper().str.extract(r"^([A-Z0-9]+)")[0]
    por_moneda = {}
    con_idx = con_valor[~con_valor["_idx"].str.upper().isin(SIN_INDICE_PLANTILLA)].assign(_ccy=lambda d: d["PK2"].map(ccy_pos), _fam=lambda d: d["PK2"].map(fam_pos))
    for (fam, ccy), g in con_idx.dropna(subset=["_ccy", "_fam"]).groupby(["_fam", "_ccy"]):
        moda = g.groupby(["_idx", "Spread_Anual"], dropna=False).size().sort_values(ascending=False).index[0]
        por_moneda[(fam, ccy)] = (moda[0], moda[1], int(g["Dias"].median()) if g["Dias"].notna().any() else 1)
    filas = []
    for _, p in sin.iterrows():
        pk2, nombre, ccy = str(p["PK2"]), str(p.get("Name_Instrumento", "")), str(p.get("Risk_Currency", "")).upper()
        fam = (re.match(r"^([A-Z0-9]+)", nombre.strip().upper()) or [None, ""])[1]
        idx, spread, dias, com, sug = None, None, None, "", ""
        m = _NOMBRE_TASA.match(nombre.strip().upper())
        if pk2 in por_pk2:
            r = por_pk2[pk2]
            idx, spread, dias, sug = r["_idx"], r["Spread_Anual"], r["Dias"], "COPIA_PK2"
            com = f"copiado de la fila del fondo {int(r['ID_Fund']) if pd.notna(r['ID_Fund']) else 'global'}: confirmar"
        elif m and _fecha_nombre(m.group(1)) is not None:
            venc = _fecha_nombre(m.group(1))
            idx, spread, dias, sug = "N.A.", round(float(m.group(2)) / 100 * 12, 6), max((venc - settle).days, 0), "NOMBRE"
            com = f"del nombre: tasa mensual {m.group(2)} % × 12, vence {venc:%Y-%m-%d}: confirmar"
        elif (fam, ccy) in por_moneda:
            idx, spread, dias = por_moneda[(fam, ccy)]
            sug, com = "POR_MONEDA", f"índice y spread más usados en cajas para {fam} en {ccy}: revisar"
        else:
            com = f"{nombre}: sin referencia (cuenta corriente / fondo money market): completar o dejar vacío = yield 0"
        filas.append(dict(ID_Fund=int(p["ID_Fund"]), PK2=pk2, Indice_Referencia=idx, Spread_Anual=spread, Dias=dias, Comentario=com, _Sugerencia=sug,
                          _Fondo=p.get("Fondo", ""), _Nombre=nombre, _Bucket=p.get("Bucket", ""), _Risk_Currency=ccy, _TotalMVal=p.get("TotalMVal")))
    out = pd.DataFrame(filas, columns=cols)
    return out.sort_values("_TotalMVal", key=lambda s: s.abs(), ascending=False).reset_index(drop=True)
