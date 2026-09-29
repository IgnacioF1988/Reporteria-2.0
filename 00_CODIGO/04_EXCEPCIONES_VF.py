# =============================================================================
# 04_EXCEPCIONES.py — Métricas desde flujos provistos manualmente
# =============================================================================
# QUÉ SON LAS EXCEPCIONES
#   Papeles cuya TD la entrega directamente el analista/PM en EXCEPCIONES.xlsx,
#   una pestaña por PK2. Cubren tres situaciones:
#     · estructuras únicas que ningún proveedor modela (loans, distress, PIK)
#     · papeles ausentes del bond_schedule.jsonl
#     · papeles donde al PM NO le convence la yield del proveedor y prefiere
#       usar sus propios flujos
#
# PRECEDENCIA
#   EXCEPCIONES PISA a todos los proveedores. Si un PK2 está acá, su Yield y
#   Duration salen de esta tabla aunque JPM/RA/BBG/jsonl tengan valor.
#   (Por eso en la cascada de pipeline_config va primero.)
#
# ESCALA DE LOS FLUJOS
#   Los flujos vienen normalizados a un nominal de FACE_EXCEPCIONES (1.000.000).
#   Se normaliza así porque un mismo PK2 se reparte entre fondos con Qty
#   distinto: se guarda una sola tabla y cada posición la escala con su Q_real.
#       flujo_posicion = Flujo × (Q_real / 1.000.000)
#   Q_real = Qty × sQ  (sQ = escala de unidad del papel; USD normalmente 1,
#   pero los locales traen otras escalas — de ahí que no se pueda omitir).
#
# FLUJO INICIAL
#   NO viene en la tabla: se toma del CUBO por (PK2, ID_Fund), porque cada
#   fondo puede valorizar el mismo papel con distinta fuente de precios.
#       FI = -(P_ef × Q_real) - AI_local        [precio sucio, moneda local]
#   Los escalares (sP, sQ) y el FX se identifican calzando contra MVBook,
#   misma lógica que 02_LIMPIEZA / 03_CALCULOS.
#
# BALANCESHEET
#   El foco son los ACTIVOS. Las filas 'Liability' se reportan aparte y no
#   entran al cálculo (su signo es opuesto y hoy no se modelan).
#
# OUTPUT: EXCEPCIONES_{FECHA}.xlsx
#   - metricas     : Yield / Duration por posición (fondo × PK2)
#   - td_detalle   : flujos escalados usados en cada cálculo
#   - sin_calcular : PK2 de la planilla que no se pudieron resolver
#   - RESUMEN
# =============================================================================

import os
import sys
import warnings
import numpy as np
import pandas as pd
from scipy.optimize import brentq
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

_sd = os.path.dirname(os.path.abspath(__file__)) if '__file__' in dir() else os.getcwd()
for p in (_sd, os.getcwd()):
    if p not in sys.path: sys.path.insert(0, p)

from pipeline_config import (
    FECHA, excepciones_paths, CUBO_PATH, BD_FUNDS_PATH, OUT_EXCEPCIONES,
    OUT_UNIVERSO, FACE_EXCEPCIONES, INDEXADAS, NAVY, RED, AMBER_F,
    CANDIDATOS_ESCALA, UMBRAL_PCT, UMBRAL_TC, RATIO_MIN, RATIO_MAX,
    RISK_CCY_TO_FX, INSTRUMENTOS_BEE, BEE_CONN, ARCHIVO_PARIDS,
    MAX_DIAS_ATRAS_PAR, FONT_NAME, FONDOS_BASE_CLP,
)

settle_dt = pd.to_datetime(FECHA)
print("=" * 70)
print(f"04_EXCEPCIONES | FECHA={FECHA}")
print("=" * 70)


# ── Helpers ──────────────────────────────────────────────────────────────────
def limpiar_txt(s):
    return (s.fillna('').astype(str).str.strip()
            .replace({'nan': '', 'None': '', 'NaN': '', '<NA>': ''}))

def fmt_ws(ws, hexc, freeze='B2'):
    fill = PatternFill('solid', fgColor=hexc)
    font = Font(color='FFFFFF', bold=True, size=9, name='Calibri')
    for c in ws[1]:
        c.fill = fill; c.font = font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if freeze: ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 26
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 9), 40)

def xirr_calc(flujos, fechas):
    if len(flujos) < 2: return np.nan
    t0 = min(pd.to_datetime(f) for f in fechas)
    days = [(pd.to_datetime(f) - t0).days / 365.25 for f in fechas]
    def npv(r): return sum(c / (1 + r) ** d for c, d in zip(flujos, days))
    try:    return brentq(npv, -0.9999, 100.0, maxiter=200)
    except Exception: return np.nan

def duration_calc(ff, fe, s, x):
    if pd.isna(x) or not ff: return np.nan, np.nan
    s = pd.to_datetime(s)
    t = [(pd.to_datetime(f) - s).days / 365.25 for f in fe]
    pv = [cf / (1 + x) ** ti for cf, ti in zip(ff, t)]
    tot = sum(pv)
    if tot <= 0: return np.nan, np.nan
    mac = sum(ti * p for ti, p in zip(t, pv)) / tot
    return mac, mac / (1 + x)

def evaluar_escala(p, of, factor, mv, fx_val):
    """Idéntico a evaluar() de 02_LIMPIEZA_VF: 3 checks contra MVBook."""
    res = []
    for sP, sQ in CANDIDATOS_ESCALA:
        p_ef, of_ef = p * sP, of * sQ
        pq   = p_ef * of_ef * factor / fx_val
        diff = abs(pq - mv) / abs(mv) if mv != 0 else np.nan
        tc   = (p_ef * of_ef * factor) / abs(mv) if mv != 0 else np.nan
        sanity = True if fx_val == 1.0 else (
            pd.notna(tc) and abs(tc - fx_val) / fx_val <= UMBRAL_TC)
        ratio = (of_ef * factor / fx_val) / abs(mv) if (fx_val != 0 and mv != 0) else np.nan
        ratio_ok = pd.notna(ratio) and (RATIO_MIN <= ratio <= RATIO_MAX)
        res.append({'sP': sP, 'sQ': sQ, 'diff_pct': diff, 'ratio': ratio,
                    'ratio_ok': bool(ratio_ok), 'sanity_ok': bool(sanity),
                    'all_ok': pd.notna(diff) and diff <= UMBRAL_PCT and sanity and ratio_ok})
    return res


# =============================================================================
# [1] LEER EXCEPCIONES — varios archivos, una pestaña por PK2
# =============================================================================
# Se admiten VARIOS archivos con la misma estructura (EXCEPCIONES_USD.xlsx,
# EXCEPCIONES_MLDL.xlsx, EXCEPCIONES_MRCLP.xlsx…), solo para tenerlo ordenado
# por grupo de fondos. NO hay tratamiento distinto por archivo: se leen todos
# y se unen. Un PK2 debe estar en UNO solo.
archivos = excepciones_paths()
if not archivos:
    print(f"   ERROR: no se encontró ningún EXCEPCIONES*.xlsx en la carpeta MANUALES")
    sys.exit(1)
print(f"\n[1] Leyendo {len(archivos)} archivo(s) de excepciones...")

flujos_pk2, liab_rows, aviso, origen_pk2, duplicados = {}, [], [], {}, []
for ruta in archivos:
    xl = pd.ExcelFile(ruta)
    n_pk2_arch = 0
    for sh in xl.sheet_names:
        df = pd.read_excel(ruta, sheet_name=sh)
        df.columns = [str(c).strip() for c in df.columns]
        if not {'Fecha', 'Flujo'}.issubset(df.columns):
            aviso.append({'Archivo': os.path.basename(ruta), 'PK2': sh,
                          'Motivo': 'faltan columnas Fecha/Flujo'}); continue
        pk2 = (str(df['PK2'].dropna().iloc[0]).strip()
               if 'PK2' in df.columns and df['PK2'].notna().any() else sh.strip())

        if pk2 in flujos_pk2:      # mismo PK2 en dos archivos → se avisa
            duplicados.append({'PK2': pk2, 'Archivo_usado': origen_pk2[pk2],
                               'Archivo_ignorado': os.path.basename(ruta)})
            continue

        df['Fecha'] = pd.to_datetime(df['Fecha'], errors='coerce')
        df['Flujo'] = pd.to_numeric(df['Flujo'], errors='coerce')
        df = df.dropna(subset=['Fecha', 'Flujo'])
        if 'BalanceSheet' in df.columns:
            bs = df['BalanceSheet'].astype(str).str.strip().str.upper()
            n_liab = int((bs == 'LIABILITY').sum())
            if n_liab:
                liab_rows.append({'PK2': pk2, 'filas_liability': n_liab})
            df = df[bs != 'LIABILITY']
        flujos_pk2[pk2] = df[df['Fecha'] > settle_dt].sort_values('Fecha')
        origen_pk2[pk2] = os.path.basename(ruta)
        n_pk2_arch += 1
    print(f"   {os.path.basename(ruta):32} {len(xl.sheet_names):>3} pestañas → {n_pk2_arch:>3} PK2")

print(f"   TOTAL PK2 con flujos: {len(flujos_pk2)}")
print(f"   Flujos futuros totales: {sum(len(v) for v in flujos_pk2.values())}")
if duplicados:
    print(f"   [WARN] {len(duplicados)} PK2 repetidos entre archivos "
          f"(se usa la primera aparición): {[d['PK2'] for d in duplicados][:5]}")
if liab_rows:
    print(f"   [INFO] {len(liab_rows)} PK2 con filas Liability (excluidas: el foco son activos)")

# =============================================================================
# [2] UNIVERSO — posiciones (fondo × PK2) de esos papeles
# =============================================================================
print(f"\n[2] Ubicando posiciones en el universo...")
uni = pd.concat([pd.read_excel(OUT_UNIVERSO, sheet_name=s, engine='openpyxl')
                 for s in ('con_ISIN', 'sin_ISIN')], ignore_index=True)
uni['PK2'] = limpiar_txt(uni['PK2'])
pos = uni[uni['PK2'].isin(flujos_pk2.keys())].copy()
print(f"   Posiciones a calcular: {len(pos)} | PK2: {pos['PK2'].nunique()} de {len(flujos_pk2)}")
faltan_uni = sorted(set(flujos_pk2) - set(pos['PK2']))
if faltan_uni:
    print(f"   [WARN] {len(faltan_uni)} PK2 de EXCEPCIONES no están en el universo: {faltan_uni[:5]}")

# =============================================================================
# [3] CUBO
# =============================================================================
print(f"\n[3] Cruzando con CUBO...")
cubo = pd.read_excel(CUBO_PATH, engine='openpyxl')
cubo['PK2'] = limpiar_txt(cubo['PK2'])
cc = [c for c in ['ID_Fund', 'PK2', 'Qty', 'OriginalFace', 'Factor', 'AI',
                  'LocalPrice', 'MVBook', 'TotalMVal'] if c in cubo.columns]
cubo_slim = cubo[cc].drop_duplicates(['ID_Fund', 'PK2'])

if 'ID_Fund' not in pos.columns:
    bdf = pd.read_excel(BD_FUNDS_PATH, engine='openpyxl')
    cs = next((c for c in bdf.columns if 'shortname' in c.lower().replace('_', '')), None)
    pos['ID_Fund'] = pos['FundShortName'].astype(str).str.strip().map(
        dict(zip(bdf[cs].astype(str).str.strip(), bdf['ID_Fund'])))

# El universo YA arrastra las columnas de valor del CUBO (LocalPrice…TotalMVal).
# Se eliminan de pos para que el CUBO sea la fuente autoritativa y no queden
# columnas duplicadas con sufijo.
_dup = [c for c in cubo_slim.columns if c not in ('ID_Fund', 'PK2') and c in pos.columns]
if _dup:
    print(f"   (columnas tomadas del CUBO, no del universo: {_dup})")
    pos = pos.drop(columns=_dup)

pos = pos.merge(cubo_slim, on=['ID_Fund', 'PK2'], how='left')
print(f"   Posiciones: {len(pos)} | sin MVBook: {pos['MVBook'].isna().sum() if 'MVBook' in pos else 'n/a'}")

# =============================================================================
# [4] FX (beemining + paridades) — para llevar AI a moneda local
# =============================================================================
print(f"\n[4] Cargando FX...")
fx_bee, fx_par = {'USDUSD': 1.0}, {'USDUSD': 1.0}
try:
    import pyodbc, pandas.io.sql as psql
    conn = pyodbc.connect(BEE_CONN)
    ph = ",".join(f"'{c}'" for c in INSTRUMENTOS_BEE)
    sql = (f"SELECT instrumentcode,instrumentvalue,daydate FROM (SELECT instrumentcode,"
           f"instrumentvalue,daydate,ROW_NUMBER() OVER (PARTITION BY instrumentcode "
           f"ORDER BY daydate DESC) rn FROM [DW_MONEDA].[dbo].[TBL_RENTABILIDADES_DW] "
           f"WHERE daydate<='{FECHA}' AND instrumentcode IN ({ph})) t WHERE rn=1")
    for _, r in psql.read_sql(sql, conn).iterrows():
        c = str(r['instrumentcode']); v = float(r['instrumentvalue'])
        if v > 0: fx_bee["USDCLP" if c == "CLFXDOOB_sindesf" else c.replace(" Curncy", "")] = v
except Exception as e:
    print(f"   [WARN] beemining: {e}")

try:
    frames = []
    for hoja in ["Data Paridad NY", "Data Paridad LDN", "Data EUR|USD OBS"]:
        dp = pd.read_excel(ARCHIVO_PARIDS, sheet_name=hoja, engine='openpyxl')
        raw = dp['Date']
        if pd.api.types.is_numeric_dtype(raw):
            from datetime import date as _d
            dp['Date'] = pd.to_datetime(raw.apply(
                lambda x: _d.fromordinal(int(x) + 693594) if pd.notna(x) and x > 0 else None), errors='coerce')
        else:
            p_ = pd.to_datetime(raw, errors='coerce')
            if p_.isna().sum() > len(p_) * 0.5: p_ = pd.to_datetime(raw, errors='coerce', dayfirst=True)
            dp['Date'] = p_
        dp = dp[(dp['Date'].notna()) & (dp['Date'] <= settle_dt) &
                (dp['Date'] >= settle_dt - pd.Timedelta(days=MAX_DIAS_ATRAS_PAR))]
        dp = dp.rename(columns={'Unnamed: 5': 'Nomenclatura'}); dp['Source'] = hoja
        frames.append(dp)
    dpar = pd.concat(frames, ignore_index=True).sort_values('Date', ascending=False)
    dpar = dpar.drop_duplicates(subset=['Nomenclatura', 'Source'])
    for _, r in dpar.iterrows():
        n = str(r.get('Nomenclatura', '')).strip()
        try: v = float(r['Price'])
        except Exception: continue
        if n and v > 0: fx_par.setdefault(n, v)
    for a, b in [('CLFUSD', 'USDCLF'), ('EURUSD', 'USDEUR'), ('GBPUSD', 'USDGBP'),
                 ('UVR CurncyUSD', 'USDUVR')]:
        if fx_par.get(a, 0) > 0 and b not in fx_par: fx_par[b] = 1 / fx_par[a]
except Exception as e:
    print(f"   [WARN] paridades: {e}")
print(f"   FX bee={len(fx_bee)} par={len(fx_par)}")

def fx_candidatos(ccy):
    c = str(ccy).strip().upper()
    if c == 'USD': return [('USDUSD', 1.0)]
    base = RISK_CCY_TO_FX.get(c, f'USD{c}')
    out, seen = [], set()
    for lbl, d in (('bee', fx_bee), ('par', fx_par)):
        for k, v in d.items():
            if k.startswith(base) and pd.notna(v) and v > 0 and round(v, 6) not in seen:
                out.append((f'{lbl}_{k}', v)); seen.add(round(v, 6))
    return out

# =============================================================================
# [5] CALCULAR Yield / Duration por posición
# =============================================================================
print(f"\n[5] Calculando...")
res, td_out, sin_calc = [], [], []

for _, r in pos.iterrows():
    pk2 = r['PK2']; fondo = r.get('FundShortName', ''); fid = r.get('ID_Fund', np.nan)
    rc = str(r.get('Risk_Currency', '')).strip().upper()
    lp = pd.to_numeric(r.get('LocalPrice'), errors='coerce')
    of = pd.to_numeric(r.get('OriginalFace'), errors='coerce')
    qty = pd.to_numeric(r.get('Qty'), errors='coerce')
    fac = pd.to_numeric(r.get('Factor'), errors='coerce'); fac = fac if pd.notna(fac) else 1.0
    mvb = pd.to_numeric(r.get('MVBook'), errors='coerce')
    ai = pd.to_numeric(r.get('AI'), errors='coerce'); ai = ai if pd.notna(ai) else 0.0

    td = flujos_pk2.get(pk2)
    if td is None or td.empty:
        sin_calc.append({'Fondo': fondo, 'PK2': pk2, 'Motivo': 'Sin flujos futuros'}); continue
    if pd.isna(lp) or pd.isna(mvb) or mvb == 0 or pd.isna(qty) or qty == 0:
        sin_calc.append({'Fondo': fondo, 'PK2': pk2,
                         'Motivo': 'Sin LocalPrice/MVBook/Qty en CUBO'}); continue

    # ── MONEDA DEL FONDO ─────────────────────────────────────────────────────
    # MVBook y AI vienen EN LA MONEDA DEL FONDO, no del papel: los fondos de
    # FONDOS_BASE_CLP contabilizan en CLP, el resto en USD.
    # En vez de convertir MVBook con UN USDCLP fijo, se generan CANDIDATOS de
    # "unidades del papel por unidad de moneda del fondo" y se deja que el
    # calce elija — misma filosofía que el FX del papel (buscar, no asumir).
    # Motivo: el USDCLP contable del fondo puede diferir del de mercado
    # (caso real: 930.86 contable vs 924.78 de beemining = 0.66% de desvío,
    # suficiente para romper el umbral de 0.5% y caer a FALLBACK).
    base_clp = (int(fid) in FONDOS_BASE_CLP) if pd.notna(fid) else False

    def candidatos_fondo():
        """(label, fx) con fx = unidades del papel por unidad de moneda del fondo."""
        base = fx_candidatos(rc)
        if not base_clp:
            return base
        out = []
        vistos = set()
        for lbl_u, u in (list(('bee_USDCLP', fx_bee[k]) for k in ['USDCLP'] if k in fx_bee)
                         + list(('par_USDCLP', fx_par[k]) for k in ['USDCLP'] if k in fx_par)):
            if not u or u <= 0: continue
            for lbl, v in base:
                fx = v / u                      # papel→USD ÷ (CLP por USD) = papel por CLP
                if fx > 0 and round(fx, 10) not in vistos:
                    out.append((f'{lbl}/{lbl_u}', fx)); vistos.add(round(fx, 10))
        return out or base

    cands = candidatos_fondo()

    # Escalares + FX calzando contra MVBook (en moneda del fondo, sin convertir)
    of_base = of if pd.notna(of) and of != 0 else qty
    sP = sQ = np.nan; fxv, fxf, flag = np.nan, 'SIN_FX', 'FALLBACK'
    for lbl, v in cands:
        ok = next((i for i in evaluar_escala(lp, of_base, fac, mvb, v) if i['all_ok']), None)
        if ok:
            sP, sQ, fxv, fxf, flag = ok['sP'], ok['sQ'], v, lbl, 'OK'; break
    if flag != 'OK':
        # FALLBACK: se prefiere SIEMPRE un candidato que pase el check de ratio.
        # Sin esto, escalares que dan el MISMO P×Q (p.ej. (0.001,1000) y (1,1))
        # empatan en diff y gana el primero de la lista, que puede inflar
        # Q_real 1000x y por lo tanto los flujos. El ratio los distingue:
        # (0.001,1000)→996.76 (fuera de rango) vs (1,1)→0.997 (correcto).
        mejor = None
        for lbl, v in cands:
            for i in evaluar_escala(lp, of_base, fac, mvb, v):
                if pd.isna(i['diff_pct']): continue
                clave = (not i['ratio_ok'], i['diff_pct'])   # ratio_ok manda sobre diff
                if mejor is None or clave < mejor[0]:
                    mejor = (clave, i, lbl, v)
        if mejor is None:
            sin_calc.append({'Fondo': fondo, 'PK2': pk2, 'Motivo': f'Sin FX para {rc}'}); continue
        _, i, lbl, v = mejor
        sP, sQ, fxv, fxf = i['sP'], i['sQ'], v, f'FALLBACK_{lbl}'

    # Q_real = Qty × sQ  (nominal vigente en unidades del papel)
    q_real   = qty * sQ
    p_ef     = lp * sP
    escala   = q_real / FACE_EXCEPCIONES     # flujos vienen base 1.000.000
    # AI viene en moneda del FONDO; fxv son unidades del papel por unidad de
    # moneda del fondo, así que la conversión es directa.
    ai_local = ai * fxv
    fi_local = -(p_ef * q_real) - ai_local

    flujos = (td['Flujo'] * escala).tolist()
    fechas = td['Fecha'].tolist()
    if all(abs(f) < 1e-9 for f in flujos):
        sin_calc.append({'Fondo': fondo, 'PK2': pk2,
                         'Motivo': 'Todos los flujos futuros son 0 (PIK sin pagos) — sin TIR'}); continue

    xef = xirr_calc([fi_local] + flujos, [settle_dt] + fechas)
    mac, mod = duration_calc(flujos, fechas, settle_dt, xef)
    es_idx = rc in INDEXADAS

    res.append({
        'Fondo': fondo, 'ID_Fund': fid, 'PK2': pk2,
        'Instrumento': r.get('Name_Instrumento', ''), 'ISIN': r.get('ISIN', ''),
        'Moneda': rc, 'Metric_Currency': rc, 'Es_Indexado': es_idx,
        'Requiere_Breakeven': es_idx, 'Hedge_Currency': r.get('Hedge_Currency', ''),
        'sP': sP, 'sQ': sQ, 'Escalar_flag': flag, 'FX_usado': fxv, 'FX_fuente': fxf,
        'Qty': qty, 'Q_real': q_real, 'P_ef_pct': p_ef * 100,
        'AI_cubo': ai, 'AI_local': ai_local, 'FI_local': fi_local,
        'MVBook': mvb, 'Base_fondo': 'CLP' if base_clp else 'USD',
        'N_flujos': len(flujos),
        'Primer_flujo': min(fechas), 'Ultimo_flujo': max(fechas),
        'Yield_efectiva': xef, 'MacDur': mac, 'ModDur': mod,
        'Fuente': 'EXCEPCIONES', 'Archivo_origen': origen_pk2.get(pk2, ''),
    })
    for f_, v_ in zip(fechas, flujos):
        td_out.append({'PK2': pk2, 'ID_Fund': fid, 'Fondo': fondo,
                       'Fecha': f_, 'Flujo_base1MM': v_ / escala if escala else np.nan,
                       'Flujo_escalado': v_, 'Moneda': rc})

df_res = pd.DataFrame(res)
df_td  = pd.DataFrame(td_out)
df_sin = pd.DataFrame(sin_calc)
ok = int(df_res['Yield_efectiva'].notna().sum()) if len(df_res) else 0
print(f"   Posiciones con Yield: {ok}/{len(pos)}")
print(f"   Sin calcular        : {len(df_sin)}")

# =============================================================================
# [6] GUARDAR
# =============================================================================
rr = []
def add(m, v='', c=''): rr.append({'Concepto': m, 'Valor': v, 'Comentario': c})
add('── ENTRADA ──')
add('Archivos de excepciones', len(archivos), ', '.join(os.path.basename(a) for a in archivos))
for a in archivos:
    n = sum(1 for k, v in origen_pk2.items() if v == os.path.basename(a))
    add(f'  {os.path.basename(a)}', n, 'PK2')
add('PK2 totales', len(flujos_pk2))
if duplicados:
    add('PK2 repetidos entre archivos', len(duplicados), 'se usa la primera aparición')
add('Posiciones (fondo × PK2)', len(pos))
add('')
add('── RESULTADO ──')
add('Posiciones con Yield/Duration', ok, 'EXCEPCIONES pisa a todos los proveedores')
add('Sin calcular', len(df_sin))
if len(df_sin):
    for k, v in df_sin['Motivo'].value_counts().items():
        add(f'  {k}', int(v))
add('')
add('── ESCALA ──')
add('Nominal de normalización', f'{FACE_EXCEPCIONES:,.0f}',
    'flujo_posicion = Flujo × Q_real / 1.000.000 ; Q_real = Qty × sQ')
if len(df_res):
    add('Escalar+FX calzado con MVBook', int((df_res['Escalar_flag'] == 'OK').sum()))
    add('')
    add('── MONEDA ──')
    add('Indexados (→ breakeven)', int(df_res['Es_Indexado'].sum()))
    for k, v in df_res['Metric_Currency'].value_counts().items():
        add(f'  {k}', int(v))
if liab_rows:
    add('')
    add('── BALANCESHEET ──')
    add('PK2 con filas Liability', len(liab_rows), 'excluidas: el foco son activos')

with pd.ExcelWriter(OUT_EXCEPCIONES, engine='openpyxl') as w:
    pd.DataFrame(rr).to_excel(w, sheet_name='RESUMEN', index=False)
    fmt_ws(w.sheets['RESUMEN'], '2E4057', freeze='A2')
    (df_res if len(df_res) else pd.DataFrame({'info': ['sin resultados']})
     ).to_excel(w, sheet_name='metricas', index=False)
    fmt_ws(w.sheets['metricas'], NAVY)
    if len(df_td):
        df_td.to_excel(w, sheet_name='td_detalle', index=False)
        fmt_ws(w.sheets['td_detalle'], '2E5C9A')
    if len(df_sin):
        df_sin.to_excel(w, sheet_name='sin_calcular', index=False)
        fmt_ws(w.sheets['sin_calcular'], RED)
    if duplicados:
        pd.DataFrame(duplicados).to_excel(w, sheet_name='PK2_duplicados', index=False)
        fmt_ws(w.sheets['PK2_duplicados'], RED)

print(f"\n{'='*70}")
print(f"  Con Yield: {ok} | Sin calcular: {len(df_sin)}")
print(f"  Output: {OUT_EXCEPCIONES}")
print(f"{'='*70}")