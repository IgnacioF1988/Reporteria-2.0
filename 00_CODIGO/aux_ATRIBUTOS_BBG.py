# =============================================================================
# 04_ATRIBUTOS.py — Atributos de indexación y moneda por PK2
# =============================================================================
# PROPÓSITO
#   La cobertura (MAPEO_METRICAS) entrega Yield y Duration, pero NO dice en qué
#   moneda están. Muchas vienen en moneda indexada (UF, UDI, UVR, IPCA...) y
#   son tasas REALES: requieren breakeven para pasarlas a moneda local nominal.
#   Este script determina, por PK2:
#     · Index_Type  (NOMINAL / UF / UDI / UVR / IPCA / UI / BONCER / VAC / CDI...)
#     · Metric_Currency: moneda en la que viene la métrica de CADA fuente
#     · Requiere_Breakeven: bandera para el paso siguiente
#
# MONEDA POR FUENTE (principio: la yield viene en la moneda de los flujos/bono)
#   BBG YAS  → moneda del bono. Si CLF/UDI/UVR/... → yield REAL (indexada).
#              Señales: INFLATION_LINKED_INDICATOR='Y', CPN_TYP='FLOATING'+RESET_IDX, CRNCY
#   BBG CSHF → Flow_Currency (misma lógica; la XIRR sale en esa moneda)
#   RA       → columna MONEDA literal ('UF' / 'CLP' / 'USD')  [de RA_TIR.xlsx]
#   JPM CEMBI→ USD (nominal, sin breakeven)
#   JPM GBI  → local nominal (sin breakeven, salvo excepción marcada)
#   PROP     → moneda de los flujos del bond_schedule.jsonl (Geneva)
#
# INPUTS
#   MAPEO_METRICAS_{FECHA}.xlsx  (con_ISIN / sin_ISIN)   — universo + fuente elegida
#   RA_TIR.xlsx  (hoja del mes)                          — MONEDA por nemotécnico
#   [BBG en vivo]  INFLATION_LINKED_INDICATOR, CPN_TYP, RESET_IDX, CRNCY,
#                  MATURITY, CPN, CPN_FREQ, DAY_CNT   (para indexación + PROP)
#
# OUTPUT
#   ATRIBUTOS_{FECHA}.xlsx
#     - atributos : PK2 + Index_Type + Metric_Currency + Requiere_Breakeven + attrs PROP
#     - indexados : subconjunto que necesita breakeven (para 05_BREAKEVEN)
#     - nuevos_idx: RESET_IDX no mapeados (revisar INDEX_MAP)
#     - RESUMEN
# =============================================================================

import os
import sys
import time
import warnings
import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

_sd = os.path.dirname(os.path.abspath(__file__)) if '__file__' in dir() else os.getcwd()
for p in (_sd, os.getcwd()):
    if p not in sys.path: sys.path.insert(0, p)

from pipeline_config import (FECHA, DIR_OUTPUTS, DIR_MANUALES, OUT_METRICAS,
                             INDEX_MAP, INFLATION_MAP, RA_TIR_PATH, NAVY, RED, FONT_NAME)

# ============================================================
# CONFIGURACIÓN
# ============================================================
CONSULTAR_BBG = True
HOJA_RA       = "jul26"  # ajustar según el mes

FNAME_METRICAS = OUT_METRICAS
FNAME_RA       = RA_TIR_PATH
FNAME_OUT      = os.path.join(DIR_OUTPUTS, f"ATRIBUTOS_{FECHA}.xlsx")

# Monedas que son indexadas (yield viene REAL → breakeven)
INDEXADAS = {'UF', 'UDI', 'UVR', 'IPCA', 'UI', 'BONCER', 'VAC'}
# Risk_Currency → Index_Type cuando la moneda es inequívoca (fallback)
CCY_INDEX_DIRECTO = {'CLF': 'UF'}
# Risk_Currency ambiguas: pueden ser nominal O indexado → decide BBG
CCY_AMBIGUA = {'MXN', 'BRL', 'COP', 'ARS', 'UYU', 'PEN', 'MXV', 'CLP'}

BBG_BATCH = 100
# ============================================================

print("=" * 70)
print(f"aux_ATRIBUTOS | FECHA={FECHA} | BBG={'ON' if CONSULTAR_BBG else 'OFF'}")
print("=" * 70)


def norm(x): return str(x).strip()

def limpiar_txt(s):
    return (s.fillna('').astype(str).str.strip()
            .replace({'nan': '', 'None': '', 'NaN': '', '<NA>': ''}))

def pick_col(df, *c):
    for x in c:
        if x in df.columns: return x
    return None

def fmt_ws(ws, hexc, freeze='B2'):
    fill = PatternFill('solid', fgColor=hexc)
    font = Font(color='FFFFFF', bold=True, size=9, name='Calibri')
    aln  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c in ws[1]:
        c.fill = fill; c.font = font; c.alignment = aln
    if freeze: ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 26
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 9), 40)


# =============================================================================
# [1] UNIVERSO
# =============================================================================
print(f"\n[1] Leyendo {os.path.basename(FNAME_METRICAS)}...")
con = pd.read_excel(FNAME_METRICAS, sheet_name='con_ISIN', engine='openpyxl')
sin = pd.read_excel(FNAME_METRICAS, sheet_name='sin_ISIN', engine='openpyxl')
df  = pd.concat([con, sin], ignore_index=True)

for c in ('ISIN', 'Name_Instrumento', 'Risk_Currency', 'Risk_Country',
          'RA_Moneda', 'Fuente_Yield', 'serie_hermanos'):
    if c in df.columns:
        df[c] = limpiar_txt(df[c])

instr = df.drop_duplicates('PK2').copy()
print(f"   Instrumentos únicos (PK2): {len(instr):,}")
isins = sorted(set(instr.loc[instr['ISIN'] != '', 'ISIN']))
print(f"   ISIN únicos: {len(isins):,}")

hermanos_map = {}
for _, r in instr.iterrows():
    h = r.get('serie_hermanos', '')
    if h:
        hermanos_map[r['ISIN']] = [x.strip() for x in h.split(',') if x.strip()]


# =============================================================================
# [2] BBG — campos de indexación + atributos PROP
# =============================================================================
bbg = {}   # isin -> dict de campos crudos

if CONSULTAR_BBG and isins:
    from xbbg import blp
    FLDS = ['INFLATION_LINKED_INDICATOR', 'CPN_TYP', 'RESET_IDX', 'CRNCY',
            'MATURITY', 'CPN', 'CPN_FREQ', 'DAY_CNT']
    print(f"\n[2] BBG — {len(isins)} ISIN × {len(FLDS)} campos...")
    t0 = time.perf_counter()

    def bdp_batch(tickers, flds):
        out = {}
        for i in range(0, len(tickers), BBG_BATCH):
            lote = tickers[i:i + BBG_BATCH]
            try:
                res = blp.bdp(tickers=lote, flds=flds)
                if res is None or res.empty: continue
                res.index   = [norm(x) for x in res.index]
                res.columns = [norm(c).upper() for c in res.columns]
                for tkr in res.index:
                    out[tkr] = {c: res.loc[tkr].get(c) for c in res.columns}
            except Exception as e:
                print(f"      [WARN] lote {i}: {e}")
        return out

    raw = bdp_batch([f"{i} Corp" for i in isins], FLDS)
    for isin in isins:
        d = raw.get(f"{isin} Corp")
        if d: bbg[isin] = d
    pend = [i for i in isins if i not in bbg]
    if pend:
        raw2 = bdp_batch([f"{i}@BGN Corp" for i in pend], FLDS)
        for isin in pend:
            d = raw2.get(f"{isin}@BGN Corp")
            if d: bbg[isin] = d
    print(f"   Con atributos BBG: {len(bbg)}/{len(isins)} | {time.perf_counter()-t0:0.1f}s")
else:
    print("\n[2] BBG OMITIDO")


# =============================================================================
# [3] RA — MONEDA por nemotécnico
# =============================================================================
print(f"\n[3] RA — {os.path.basename(FNAME_RA)}")
ra_mon = {}
if os.path.exists(FNAME_RA):
    _ra = pd.read_excel(FNAME_RA, sheet_name=HOJA_RA, engine='openpyxl')
    cols = list(_ra.columns)
    _ra = _ra.rename(columns={cols[0]: 'nm'})
    _ra['nm'] = limpiar_txt(_ra['nm']).str.upper()
    cm = pick_col(_ra, 'MONEDA')
    if cm:
        for _, r in _ra.iterrows():
            if r['nm']: ra_mon[r['nm']] = norm(r.get(cm, ''))
    print(f"   RA monedas: {len(ra_mon)} | valores: {sorted(set(ra_mon.values()))}")
else:
    print(f"   [WARN] no encontrado")


# =============================================================================
# [4] DERIVAR Index_Type + Metric_Currency  (cascada BBG → RA → fallback)
# =============================================================================
print(f"\n[4] Derivando Index_Type y moneda de métrica...")

nuevos_idx = []

def index_type_bbg(isin, country):
    """Replica la lógica de 01_BBG_EXTRACCION: FLOATING→INDEX_MAP, INFLATION→INFLATION_MAP."""
    d = bbg.get(isin)
    if not d: return None, None
    reset = norm(d.get('RESET_IDX', '')).upper()
    inf   = norm(d.get('INFLATION_LINKED_INDICATOR', '')).upper()
    cpn   = norm(d.get('CPN_TYP', '')).upper()
    crncy = norm(d.get('CRNCY', '')).upper()
    if cpn == 'FLOATING' and reset and '#N/A' not in reset:
        it = INDEX_MAP.get(reset)
        if it is None:
            it = f'NUEVO:{reset}'
            nuevos_idx.append({'ISIN': isin, 'RESET_IDX': reset, 'Risk_Country': country})
        return it, crncy
    if inf == 'Y':
        return INFLATION_MAP.get(country, f'INFLACION:{country}'), crncy
    return 'NOMINAL', crncy

rows = []
for _, r in instr.iterrows():
    isin = r['ISIN']; nm = r['Name_Instrumento'].upper()
    rc   = r['Risk_Currency'].upper(); country = r['Risk_Country'].upper()
    fuente = r.get('Fuente_Yield', '')

    it_bbg, crncy_bbg = (index_type_bbg(isin, country) if isin else (None, None))
    ra_m = ra_mon.get(nm, '')

    # Cascada de Index_Type
    origen_it = ''
    if it_bbg and not str(it_bbg).startswith('NUEVO'):
        index_type = it_bbg; origen_it = 'BBG'
    elif ra_m in INDEXADAS:
        index_type = ra_m; origen_it = 'RA'
    elif rc in CCY_INDEX_DIRECTO:
        index_type = CCY_INDEX_DIRECTO[rc]; origen_it = 'CCY'
    elif it_bbg:                        # NUEVO:xxx
        index_type = it_bbg; origen_it = 'BBG_NUEVO'
    else:
        index_type = 'NOMINAL'; origen_it = 'DEFAULT'

    # Moneda en que viene la métrica de la FUENTE ELEGIDA
    if fuente == 'RA':
        metric_ccy = ra_m or rc
    elif fuente == 'JPM':
        metric_ccy = 'USD' if r.get('JPM_Fuente', '') == 'CEMBI' else rc
    elif fuente == 'BBG':
        metric_ccy = crncy_bbg or rc
    else:
        metric_ccy = rc

    # ¿la métrica es real (indexada) → breakeven?
    es_indexado = (index_type in INDEXADAS)
    # JPM CEMBI es USD nominal aunque el papel sea local: no breakeven
    requiere_be = es_indexado and not (fuente == 'JPM' and metric_ccy == 'USD')

    # Atributos PROP (para reconstruir TD si toca)
    d = bbg.get(isin, {})
    rows.append({
        'PK2': r['PK2'], 'ISIN': isin, 'Name_Instrumento': r['Name_Instrumento'],
        'Risk_Currency': rc, 'Risk_Country': country,
        'Fuente_Yield': fuente,
        'Index_Type': index_type, 'Origen_Index_Type': origen_it,
        'RA_Moneda': ra_m, 'CRNCY_BBG': crncy_bbg or '',
        'Metric_Currency': metric_ccy,
        'Es_Indexado': es_indexado,
        'Requiere_Breakeven': requiere_be,
        # atributos PROP
        'MATURITY': d.get('MATURITY'), 'CPN': d.get('CPN'),
        'CPN_FREQ': d.get('CPN_FREQ'), 'DAY_CNT': d.get('DAY_CNT'),
        'CPN_TYP': d.get('CPN_TYP'),
    })

attrs = pd.DataFrame(rows)
indexados = attrs[attrs['Requiere_Breakeven']].copy()

print(f"   Indexados (requieren breakeven): {len(indexados)}")
print(f"   Por Index_Type:")
for k, v in attrs.loc[attrs['Es_Indexado'], 'Index_Type'].value_counts().items():
    print(f"      {k:<10}: {v}")
print(f"   Origen del Index_Type:")
for k, v in attrs['Origen_Index_Type'].value_counts().items():
    print(f"      {k:<10}: {v}")
if nuevos_idx:
    print(f"   [OJO] {len(nuevos_idx)} RESET_IDX nuevos sin mapear (ver hoja nuevos_idx)")


# =============================================================================
# [5] RESUMEN + GUARDAR
# =============================================================================
print(f"\n[5] Guardando {os.path.basename(FNAME_OUT)}...")

rows_r = []
def add(m, v='', d=''): rows_r.append({'Concepto': m, 'Valor': v, 'Detalle': d})
add('── UNIVERSO ──')
add('Instrumentos únicos (PK2)', len(attrs))
add('Con atributos BBG', len(bbg))
add('')
add('── INDEXACIÓN ──')
add('Indexados (requieren breakeven)', len(indexados),
    f"{len(indexados)/len(attrs):.1%} del universo")
for k, v in attrs.loc[attrs['Es_Indexado'], 'Index_Type'].value_counts().items():
    add(f'  {k}', int(v))
add('Nominales', int((~attrs['Es_Indexado']).sum()))
add('')
add('── MONEDA DE LA MÉTRICA (fuente elegida) ──')
for k, v in attrs['Metric_Currency'].value_counts().head(12).items():
    add(f'  {k}', int(v))
add('')
add('── CONTROL: cómo se determinó el Index_Type ──')
for k, v in attrs['Origen_Index_Type'].value_counts().items():
    add(f'  {k}', int(v))
add('RESET_IDX nuevos sin mapear', len(nuevos_idx), 'revisar INDEX_MAP en pipeline_config')
resumen = pd.DataFrame(rows_r)

with pd.ExcelWriter(FNAME_OUT, engine='openpyxl') as w:
    attrs.to_excel(w,     sheet_name='atributos',  index=False)
    fmt_ws(w.sheets['atributos'], '1F3864')
    indexados.to_excel(w, sheet_name='indexados',  index=False)
    fmt_ws(w.sheets['indexados'], '7A4F01')
    if nuevos_idx:
        pd.DataFrame(nuevos_idx).to_excel(w, sheet_name='nuevos_idx', index=False)
        fmt_ws(w.sheets['nuevos_idx'], '9C0006')
    resumen.to_excel(w,   sheet_name='RESUMEN',    index=False)
    fmt_ws(w.sheets['RESUMEN'], '2E4057', freeze='A2')

print(f"\n{'='*70}")
print(f"  Indexados: {len(indexados)} | Nominales: {int((~attrs['Es_Indexado']).sum())}")
print(f"  Output   : {FNAME_OUT}")
print(f"{'='*70}")