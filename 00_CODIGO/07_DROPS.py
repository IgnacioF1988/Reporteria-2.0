# =============================================================================
# 07_DROPS.py — Yield_Drop para papeles USD swapeados a moneda local
# =============================================================================
# A QUIÉN APLICA
#   TODOS los papeles con Hedge_Currency (originalmente USD, swapeados a
#   moneda local nominal: CLP, COP, MXN, PEN, BRL). NO a indexados (CLF/UVR):
#   esos ya se resolvieron en 06_BREAKEVEN.
#
# DOS COLUMNAS, NO UNA — ninguna pisa a la otra
#   BBG_XCCY_Yield : YAS_XCCY_FIXED_COUPON_EQUIVALENT pedido en 01_METRICAS.
#                    Es el swap real de mercado. En el universo real (120
#                    hedgeados) cubre 82; los otros 38 no lo traen.
#   Yield_Drop     : la calculadora propia de este script. Se calcula para
#                    TODOS los hedgeados —con y sin XCCY—, para poder
#                    comparar ambas y detectar un XCCY que no tenga sentido,
#                    no solo para "tapar el hueco" de los que no tienen XCCY.
#   Qué yield final usar (XCCY si existe, si no Yield_Drop, y si ninguna
#   existe, override) se decide en 08_OVERRIDES, no acá.
#
# SIN CURVA DISPONIBLE → WARNING EXPLÍCITO, no falla en silencio
#   Si la moneda de Hedge_Currency no está en CURVAS_DROPS (hoy: ARS, UYU) o
#   la curva no devolvió datos, el papel queda en 'sin_convertir' con el
#   motivo puntual, y se imprime un [WARN] por consola en el momento.
#
# MÉTODO
#   drop = local_all_in − pata_usd   (todo en DECIMAL) 
#     'ADD'    (CLP, COP, MXN): local_all_in = curva_local + curva_basis/100
#     'DIRECT' (PEN, BRL)     : local_all_in = una sola curva ya combinada
#   Yield_Drop = Yield_USD + drop@duration (interpolado al plazo de la
#   duration del bono, mismo criterio de un solo punto que en breakeven)
#
# CURVAS — BDS (bulk) + fallback BDH puntual
#   BDS trae toda la curva de un golpe (CURVE_TENOR_RATES). Si un tenor falta
#   en el bulk, se intenta reconstruir el ticker por patrón PREFIJO+AÑO de la
#   MISMA curva (ej. bulk trae CPXOSS2..CPXOSS20, falta CPXOSS1 y CPXOSS30 →
#   se arman esos tickers y se pide BDH puntual solo para ellos).
#   El fallback NO se autogenera si los tickers no siguen ese patrón simple
#   (México: MPSW16C, MPBSF10J) — ahí, si falta un tenor, se extrapola plano.
#
# INPUTS
#   CONSOLIDADO_{FECHA}.xlsx  (Detalle)  ← posiciones con Requiere_Drop
#   [BBG en vivo]  BDS CURVE_TENOR_RATES + BDH puntual de fallback
#
# OUTPUT: DROPS_{FECHA}.xlsx
#   convertidos (BBG_XCCY_Yield, Yield_Drop, Dif_vs_XCCY_bps por posición)
#   sin_convertir | curvas_drop | auditoria | RESUMEN
# =============================================================================

import os
import re
import sys
import warnings
import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

_sd = os.path.dirname(os.path.abspath(__file__)) if '__file__' in dir() else os.getcwd()
for p in (_sd, os.getcwd()):
    if p not in sys.path: sys.path.insert(0, p)

from pipeline_config import (
    FECHA, OUT_CONSOLIDADO, OUT_DROPS, CURVAS_DROPS, CURVAS_SIN_FALLBACK_AUTO,
    NAVY, RED, AMBER_F, GREEN_F, FONT_NAME,
)

# ============================================================
CONSULTAR_BBG = True
# ============================================================

settle_dt = pd.to_datetime(FECHA)
print("=" * 70)
print(f"07_DROPS | FECHA={FECHA}")
print("=" * 70)

auditoria = []
def audit(tipo, detalle, **kw):
    auditoria.append({'Tipo': tipo, 'Detalle': detalle, **kw})


def limpiar_txt(s):
    return (s.fillna('').astype(str).str.strip()
            .replace({'nan': '', 'None': '', 'NaN': '', '<NA>': ''}))

def fmt_ws(ws, hexc, freeze='B2'):
    fill = PatternFill('solid', fgColor=hexc)
    font = Font(color='FFFFFF', bold=True, size=9, name='Calibri')
    for c in ws[1]:
        c.fill = fill; c.font = font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if freeze: ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 26
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 9), 40)

def tenor_to_days(tenor):
    m = re.match(r'^(\d+)(D|W|M|Y)$', str(tenor).strip().upper())
    if not m: return np.nan
    n, u = int(m.group(1)), m.group(2)
    return {'D': 1, 'W': 7, 'M': 30.4375, 'Y': 365.25}[u] * n

def interp_rate(target_days, td, yd):
    """Interpolación lineal con extrapolación FLAT en los extremos."""
    v = ~np.isnan(yd)
    td, yd = np.asarray(td)[v], np.asarray(yd)[v]
    if len(td) == 0: return np.nan
    order = np.argsort(td); td, yd = td[order], yd[order]
    if target_days <= td[0]:  return float(yd[0])
    if target_days >= td[-1]: return float(yd[-1])
    return float(np.interp(target_days, td, yd))


# =============================================================================
# [1] CONSOLIDADO — universo que requiere drop
# =============================================================================
print(f"\n[1] Leyendo {os.path.basename(OUT_CONSOLIDADO)}...")
if not os.path.exists(OUT_CONSOLIDADO):
    print("   ERROR: falta el consolidado. Corré 05_CONSOLIDA primero."); sys.exit(1)

cons = pd.read_excel(OUT_CONSOLIDADO, sheet_name='Detalle', header=2, engine='openpyxl')
cons.columns = [str(c).strip() for c in cons.columns]
for c in ('PK2', 'Risk_Currency', 'Hedge_Currency', 'Name_Instrumento', 'FundShortName'):
    if c in cons.columns: cons[c] = limpiar_txt(cons[c])

objetivo = cons[(limpiar_txt(cons.get('Hedge_Currency', '')) != '') &
                (~cons.get('Es_Indexado', False).fillna(False).astype(bool))].copy()
print(f"   Requieren drop: {len(objetivo):,} | PK2 únicos: {objetivo['PK2'].nunique():,}")
if objetivo.empty:
    print("   Nada que convertir."); sys.exit(0)

# ── Prioridad 1: BBG_XCCY_Yield ya viene del swap real de mercado ───────────
# ── Se calcula Yield_Drop para TODOS los hedgeados ──────────────────────────
# Antes solo se calculaba para los que NO tenían BBG_XCCY_Yield. Con el
# universo completo (120 hedgeados) resultó que 38 no traen XCCY — no son un
# caso marginal — y además conviene tener el drop propio SIEMPRE, aunque haya
# XCCY, para poder comparar ambos y detectar un XCCY que no tenga sentido.
# BBG_XCCY_Yield sigue siendo la referencia de mercado; Yield_Drop es la
# calculadora propia, en columna aparte — ninguna pisa a la otra acá.
tiene_xccy = (objetivo['BBG_XCCY_Yield'].notna()
              if 'BBG_XCCY_Yield' in objetivo.columns else pd.Series(False, index=objetivo.index))
print(f"   Con BBG_XCCY_Yield: {int(tiene_xccy.sum()):,} | sin XCCY: {int((~tiene_xccy).sum()):,}")
print(f"   Se calcula Yield_Drop para los {len(objetivo):,} — con y sin XCCY")

pendientes = objetivo.copy()   # TODOS pasan por la calculadora propia

monedas_pend = sorted(pendientes['Hedge_Currency'].str.upper().unique())
print(f"   Monedas a calcular: {monedas_pend}")


# =============================================================================
# [2] CURVAS — BDS bulk + fallback BDH
# =============================================================================
curvas_cache = {}     # nombre_curva -> DataFrame [tenor, tenor_days, ticker, mid]
curvas_log   = []

if CONSULTAR_BBG and monedas_pend:
    from xbbg import blp

    def bds_bulk(index_ticker):
        ticker_bbg = f"{index_ticker} Index"    # las curvas (YCSW..., YCSW0490...) requieren " Index"
        try:
            d = blp.bds(ticker_bbg, "CURVE_TENOR_RATES", CURVE_DATE=FECHA)
        except Exception as e:
            print(f"      [WARN] BDS {ticker_bbg}: {e}")
            return pd.DataFrame()
        if d is None or d.empty: return pd.DataFrame()
        d = d.reset_index() if d.index.name else d
        d.columns = [str(c).strip().lower() for c in d.columns]
        c_ten = next((c for c in d.columns if c == 'tenor'), None)
        c_tkr = next((c for c in d.columns if 'ticker' in c), None)
        c_mid = next((c for c in d.columns if 'mid' in c), None)
        if not (c_ten and c_mid):
            return pd.DataFrame()
        out = pd.DataFrame({
            'tenor': d[c_ten].astype(str).str.strip().str.upper(),
            'ticker': d[c_tkr].astype(str).str.strip() if c_tkr else '',
            'mid': pd.to_numeric(d[c_mid], errors='coerce'),
        })
        out['tenor_days'] = out['tenor'].apply(tenor_to_days)
        return out.dropna(subset=['tenor_days', 'mid'])

    def bdh_puntual(ticker):
        try:
            d = blp.bdh(tickers=ticker, flds='PX_LAST', start_date=FECHA, end_date=FECHA)
            if d is None or d.empty: return np.nan
            return float(d.iloc[-1, 0])
        except Exception:
            return np.nan

    def completar_fallback(df_curva, prefijo, sufijo_hint):
        """Reconstruye tenors AÑO faltantes de la MISMA familia (prefijo+año),
        solo si el prefijo es de patrón numérico simple."""
        if not prefijo or prefijo in CURVAS_SIN_FALLBACK_AUTO:
            return df_curva, []
        # tickers año-puros ya presentes: prefijo + solo dígitos
        anios_presentes = set()
        for t in df_curva['ticker']:
            m = re.match(rf'^{re.escape(prefijo)}(\d+)$', str(t).split(' ')[0])
            if m: anios_presentes.add(int(m.group(1)))
        if not anios_presentes:
            return df_curva, []
        rescatados = []
        for anio in (1, 30):        # los extremos típicos que suelen faltar
            if anio in anios_presentes: continue
            ticker_full = f"{prefijo}{anio}{sufijo_hint}"
            v = bdh_puntual(ticker_full)
            if pd.notna(v):
                nueva = pd.DataFrame([{'tenor': f'{anio}Y', 'ticker': ticker_full,
                                       'mid': v, 'tenor_days': anio * 365.25}])
                df_curva = pd.concat([df_curva, nueva], ignore_index=True)
                rescatados.append(ticker_full)
        return df_curva, rescatados

    def cargar_curva(nombre, index_ticker, prefijo):
        if nombre in curvas_cache: return curvas_cache[nombre]
        d = bds_bulk(index_ticker)
        if d.empty:
            audit('CURVA_SIN_DATOS', f'{nombre} ({index_ticker})',
                  Impacto='BDS no devolvió datos')
            curvas_cache[nombre] = pd.DataFrame()
            return curvas_cache[nombre]
        # sufijo típico observado en el bulk (" BGN Curncy") para construir el fallback
        suf = ''
        for t in d['ticker']:
            if ' ' in str(t):
                suf = ' ' + str(t).split(' ', 1)[1]; break
        d, rescatados = completar_fallback(d, prefijo, suf)
        if rescatados:
            print(f"      {nombre}: fallback BDH → {rescatados}")
            curvas_log.append({'Curva': nombre, 'Tickers_fallback': ', '.join(rescatados)})
        elif prefijo in CURVAS_SIN_FALLBACK_AUTO:
            faltan_anios = sorted(set((1, 30)))
            audit('SIN_FALLBACK_AUTO', nombre,
                  Impacto=f'tickers no numéricos ({prefijo}): si falta 1Y/30Y se extrapola plano')
        curvas_cache[nombre] = d
        print(f"      {nombre:12} ({index_ticker}) {len(d):>3} puntos "
              f"[{d['tenor_days'].min():.0f}–{d['tenor_days'].max():.0f} días]" if len(d) else f"      {nombre}: vacío")
        return d

    print(f"\n[2] Cargando curvas...")
    for ccy in monedas_pend:
        cfg = CURVAS_DROPS.get(ccy)
        if not cfg:
            audit('MONEDA_SIN_CURVAS', ccy, Impacto='no está en CURVAS_DROPS — agregar tickers')
            continue
        cargar_curva(f'{ccy}_local', cfg['local_curve'], cfg.get('local_prefix'))
        if cfg['metodo'] == 'ADD':
            cargar_curva(f'{ccy}_basis', cfg['basis_curve'], cfg.get('basis_prefix'))
        cargar_curva(f'{ccy}_usd', cfg['usd_curve'], cfg.get('usd_prefix'))
else:
    print("\n[2] BBG OMITIDO (CONSULTAR_BBG=False o nada pendiente)")


# =============================================================================
# [3] CALCULAR DROP POR POSICIÓN
# =============================================================================
print(f"\n[3] Calculando drops...")
filas, sin_conv = [], []

for _, r in pendientes.iterrows():
    ccy = str(r['Hedge_Currency']).strip().upper()
    y_usd = pd.to_numeric(r.get('Yield'), errors='coerce')
    dur   = pd.to_numeric(r.get('Duration'), errors='coerce')
    xccy  = pd.to_numeric(r.get('BBG_XCCY_Yield'), errors='coerce')
    com = {'Fondo': r.get('FundShortName', ''), 'PK2': r['PK2'],
           'Instrumento': r.get('Name_Instrumento', ''), 'Hedge_Currency': ccy,
           'Yield_USD': y_usd, 'Duration_original': dur,
           'BBG_XCCY_Yield': xccy, 'Tiene_XCCY': pd.notna(xccy)}

    cfg = CURVAS_DROPS.get(ccy)
    if not cfg:
        motivo = f'{ccy}: sin curvas definidas en CURVAS_DROPS — no se pudo hedgear por este motivo'
        sin_conv.append({**com, 'Motivo': motivo})
        print(f"      [WARN] {r['PK2']} ({ccy}): {motivo}")
        continue
    if pd.isna(y_usd) or pd.isna(dur) or dur <= 0:
        sin_conv.append({**com, 'Motivo': 'Sin Yield o Duration válida'}); continue

    c_local = curvas_cache.get(f'{ccy}_local', pd.DataFrame())
    c_usd   = curvas_cache.get(f'{ccy}_usd', pd.DataFrame())
    if c_local.empty or c_usd.empty:
        motivo = f'Curva local o USD sin datos para {ccy} — no se pudo hedgear por este motivo'
        sin_conv.append({**com, 'Motivo': motivo})
        print(f"      [WARN] {r['PK2']} ({ccy}): {motivo}")
        continue

    dias = dur * 365.25
    r_local = interp_rate(dias, c_local['tenor_days'].values, c_local['mid'].values)
    r_usd   = interp_rate(dias, c_usd['tenor_days'].values, c_usd['mid'].values)
    if pd.isna(r_local) or pd.isna(r_usd):
        sin_conv.append({**com, 'Motivo': 'Interpolación sin resultado'}); continue

    if cfg['metodo'] == 'ADD':
        c_basis = curvas_cache.get(f'{ccy}_basis', pd.DataFrame())
        if c_basis.empty:
            sin_conv.append({**com, 'Motivo': f'Curva basis sin datos para {ccy}'}); continue
        r_basis = interp_rate(dias, c_basis['tenor_days'].values, c_basis['mid'].values)
        if pd.isna(r_basis):
            sin_conv.append({**com, 'Motivo': 'Interpolación basis sin resultado'}); continue
        local_all_in = r_local / 100 + r_basis / 100 / 100    # basis viene en bps
    else:
        r_basis = np.nan
        local_all_in = r_local / 100

    drop = local_all_in - r_usd / 100          # decimal (ej. 0.0817)
    drop_pp = drop * 100                        # a puntos porcentuales, escala de y_usd
    y_drop = y_usd + drop_pp                     # y_usd viene en % (ej. 20.598679), no decimal

    dif_vs_xccy = (y_drop - xccy) if pd.notna(xccy) else np.nan   # ambos en %

    filas.append({**com, 'plazo_dias': round(dias),
                  'r_local_%': r_local, 'r_basis_bps': r_basis, 'r_usd_%': r_usd,
                  'local_all_in': local_all_in, 'drop': drop, 'drop_pp': drop_pp,
                  'Yield_Drop': y_drop, 'Duration_hedged': dur,   # duration no cambia por el swap
                  'Dif_vs_XCCY_bps': (dif_vs_xccy * 100 if pd.notna(dif_vs_xccy) else np.nan),  # pp → bps
                  'Metodo': cfg['metodo']})

conv = pd.DataFrame(filas)
sinc = pd.DataFrame(sin_conv)
print(f"   Calculados con curva propia: {len(conv):,}")
print(f"   Sin convertir              : {len(sinc):,}")
if len(conv):
    print(f"   Drop medio por moneda:")
    for ccy, g in conv.groupby('Hedge_Currency'):
        print(f"      {ccy}: {g['drop'].mean()*100:+.2f}%  (n={len(g)})")

# 'conv' ya trae, por posición: BBG_XCCY_Yield (referencia de mercado, si
# existía) y Yield_Drop (calculadora propia) — columnas separadas, ninguna
# pisa a la otra. El criterio de cuál usar para la yield final del bono
# queda para 08_OVERRIDES (por defecto: XCCY si existe, si no Yield_Drop).
final = conv.copy()
if len(final) and final['Dif_vs_XCCY_bps'].notna().any():
    grandes = final[final['Dif_vs_XCCY_bps'].abs() > 50]     # > 0.50% de diferencia
    if len(grandes):
        print(f"   [OJO] {len(grandes)} papeles con diferencia XCCY vs Drop propio > 50 bps — revisar")


# =============================================================================
# [4] GUARDAR
# =============================================================================
print(f"\n[4] Guardando {os.path.basename(OUT_DROPS)}...")

rr = []
def add(m, v='', c=''): rr.append({'Concepto': m, 'Valor': v, 'Comentario': c})
add('── MÉTODO ──')
add('Fórmula', 'drop = local_all_in − pata_usd', 'interpolado al plazo de la Duration')
add('ADD (CLP, COP, MXN)', 'local + basis/100', 'basis viene en puntos básicos')
add('DIRECT (PEN, BRL)', 'una sola curva ya combinada', 'BRL usa Cupom Cambial, no SOFR')
add('')
add('── UNIVERSO ──')
add('Requieren drop (Hedge_Currency)', len(objetivo))
add('  con BBG_XCCY_Yield', int(tiene_xccy.sum()), 'referencia de mercado')
add('  sin BBG_XCCY_Yield', int((~tiene_xccy).sum()), 'dependen únicamente del Yield_Drop calculado')
add('  Yield_Drop calculado', len(conv), 'para TODOS los hedgeados, con y sin XCCY')
add('  sin convertir', len(sinc), 'no se pudo hedgear — ver hoja sin_convertir')
if len(sinc):
    for k, v in sinc['Motivo'].value_counts().items():
        add(f'    {k}', int(v))
add('')
if len(conv):
    add('── DROP MEDIO POR MONEDA ──')
    for ccy, g in conv.groupby('Hedge_Currency'):
        add(f'  {ccy}', f"{g['drop'].mean()*100:+.2f}%", f"n={len(g)}")
    add('')
    add('── XCCY vs YIELD_DROP (validación cruzada) ──')
    con_ambos = conv[conv['Tiene_XCCY']]
    if len(con_ambos):
        add('Papeles con ambos disponibles', len(con_ambos))
        add('  diferencia media', f"{con_ambos['Dif_vs_XCCY_bps'].mean():+.1f} bps")
        add('  diferencia máxima', f"{con_ambos['Dif_vs_XCCY_bps'].abs().max():.1f} bps")
        grandes = con_ambos[con_ambos['Dif_vs_XCCY_bps'].abs() > 50]
        add('  con diferencia > 50 bps', len(grandes), 'revisar caso a caso')
add('')
add('── AUDITORÍA ──')
if auditoria:
    aud = pd.DataFrame(auditoria)
    for k, v in aud['Tipo'].value_counts().items():
        add(f'  {k}', int(v))
else:
    add('Sin incidencias', 'OK')

with pd.ExcelWriter(OUT_DROPS, engine='openpyxl') as w:
    pd.DataFrame(rr).to_excel(w, sheet_name='RESUMEN', index=False)
    fmt_ws(w.sheets['RESUMEN'], '2E4057', freeze='A2')
    (final if len(final) else pd.DataFrame({'info': ['sin resultados']})).to_excel(
        w, sheet_name='convertidos', index=False)
    fmt_ws(w.sheets['convertidos'], NAVY)
    if len(sinc):
        sinc.to_excel(w, sheet_name='sin_convertir', index=False)
        fmt_ws(w.sheets['sin_convertir'], AMBER_F)
    if curvas_cache:
        with_data = {k: v for k, v in curvas_cache.items() if len(v)}
        if with_data:
            _curvas_df = pd.concat([v.assign(Curva=k) for k, v in with_data.items()],
                                   ignore_index=True)
            _curvas_df.to_excel(w, sheet_name='curvas_drop', index=False)
            fmt_ws(w.sheets['curvas_drop'], GREEN_F)
            # Respaldo independiente del pipeline, para auditoría o reconstrucción manual
            try:
                from pipeline_config import DIR_MERCADO
                _bak_path = os.path.join(DIR_MERCADO, f"CURVAS_DROPS_{FECHA}.csv")
                _curvas_df.to_csv(_bak_path, index=False, encoding='utf-8-sig')
                print(f"   Respaldo de curvas: {_bak_path}")
            except Exception as e:
                print(f"   [WARN] no se pudo guardar respaldo CSV de curvas: {e}")
    if auditoria:
        pd.DataFrame(auditoria).to_excel(w, sheet_name='auditoria', index=False)
        fmt_ws(w.sheets['auditoria'], RED)

print(f"\n{'='*70}")
print(f"  Yield_Drop calculado: {len(conv):,}  (con XCCY={int(conv['Tiene_XCCY'].sum()) if len(conv) else 0}, sin XCCY={int((~conv['Tiene_XCCY']).sum()) if len(conv) else 0})")
print(f"  Sin convertir       : {len(sinc):,}")
if auditoria:
    print(f"  [AUDITORÍA] {len(auditoria)} incidencias — ver hoja 'auditoria'")
print(f"  Output: {OUT_DROPS}")
print(f"{'='*70}")