# =============================================================================
# 05_CONSOLIDA.py — Consolidación + reporte ejecutivo
# =============================================================================
# Reúne las métricas de todas las etapas y resuelve UNA Yield y UNA Duration por
# posición (fondo × PK2), dejando explícita la moneda en que quedan.
#
# LEE (de 02_OUTPUTS\{FECHA}\, todo por la misma FECHA)
#   UNIVERSO_{f}.xlsx     con_ISIN / sin_ISIN     ← base
#   METRICAS_{f}.xlsx     con_ISIN / sin_ISIN     ← JPM / RA / BBG
#   CSHF_{f}.xlsx         resueltos              ← TD de Bloomberg
#   PROP_JSONL_{f}.xlsx   metricas               ← TD de Geneva
#   EXCEPCIONES_{f}.xlsx  metricas               ← flujos del PM
#   DEFAULTEADOS.xlsx                            ← DEF / PROPDEF
#   Si falta alguno, se omite esa fuente y se avisa.
#
# CASCADA:  EXCEPCIONES > JPM > RA > BBG > CSHF > JSONL
#   Se exige TUPLA COMPLETA (Yield y Duration). EXCEPCIONES va primero porque
#   es decisión explícita del PM y pisa a cualquier proveedor.
#
# DEFAULTEADOS: DEF y PROPDEF no muestran métrica de proveedor — se fuerzan a
#   Yield = 0 / Duration = 0.5 DESPUÉS de la cascada, para que pisen todo.
#
# MONEDA: el consolidado queda en la MONEDA DEL PAPEL. Nada se convierte acá:
#   solo se marcan Requiere_Breakeven (06) y Requiere_Drop (07).
#
# OUTPUT: CONSOLIDADO_{FECHA}.xlsx
#   Panorama          cifras clave del cierre
#   Resumen por fondo cobertura por fondo (conteo y AUM)
#   Detalle           una fila por posición, con métrica y procedencia
#   FALTANTES         el mapeo real de lo que sigue sin resolver
#   Defaulteados      los forzados por regla
#   Facturas          excluidas del universo relevante
# =============================================================================

import os
import sys
import warnings
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

_sd = os.path.dirname(os.path.abspath(__file__)) if '__file__' in dir() else os.getcwd()
for p in (_sd, os.getcwd()):
    if p not in sys.path: sys.path.insert(0, p)

from pipeline_config import (
    FECHA, CASCADA, INDEXADAS, YIELD_DEF, DURATION_DEF, MARCAS_DEF,
    DEFAULTEADOS_PATH, ISSUE_TYPE_BONO, PREFIJOS_EXCLUIDOS,
    OUT_UNIVERSO, OUT_METRICAS, OUT_CSHF, OUT_JSONL, OUT_EXCEPCIONES,
    OUT_CONSOLIDADO, NAVY, GREY_L, WHITE, GREEN, GREEN_F, RED, RED_F, AMBER_F,
    FONT_NAME,
)

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

print("=" * 70)
print(f"05_CONSOLIDA | FECHA={FECHA} | cascada: {' > '.join(CASCADA)}")
print("=" * 70)


def limpiar_txt(s):
    return (s.fillna('').astype(str).str.strip()
            .replace({'nan': '', 'None': '', 'NaN': '', '<NA>': ''}))

def leer(path, hoja):
    if not os.path.exists(path):
        print(f"   [WARN] falta {os.path.basename(path)} — se omite")
        return pd.DataFrame()
    try:
        d = pd.read_excel(path, sheet_name=hoja, engine='openpyxl')
        print(f"   {os.path.basename(path):30} [{hoja:12}] {len(d):>5} filas")
        return d
    except Exception as e:
        print(f"   [WARN] {os.path.basename(path)}/{hoja}: {e}")
        return pd.DataFrame()

def key(df):
    return (df['ID_Fund'].astype(str) if 'ID_Fund' in df.columns
            else pd.Series([''] * len(df), index=df.index)) + '|' + \
           df['PK2'].astype(str).str.strip()


# =============================================================================
# [1] BASE
# =============================================================================
print(f"\n[1] Base del universo...")
base = pd.concat([pd.read_excel(OUT_UNIVERSO, sheet_name=s, engine='openpyxl')
                  for s in ('con_ISIN', 'sin_ISIN')], ignore_index=True)
for c in ('PK2', 'ISIN', 'Risk_Currency', 'Hedge_Currency', 'Name_Instrumento'):
    if c in base.columns: base[c] = limpiar_txt(base[c])
if 'ID_Fund' not in base.columns: base['ID_Fund'] = np.nan
base['_k'] = key(base)
print(f"   Posiciones: {len(base):,} | PK2 únicos: {base['PK2'].nunique():,}")

# =============================================================================
# [2] APORTES
# =============================================================================
print(f"\n[2] Leyendo fuentes...")
df = base.copy()

# --- METRICAS: trae JPM / RA / BBG en una misma hoja ---
met = pd.concat([leer(OUT_METRICAS, 'con_ISIN'), leer(OUT_METRICAS, 'sin_ISIN')],
                ignore_index=True)
if not met.empty:
    met.columns = [str(c).strip() for c in met.columns]
    met['PK2'] = limpiar_txt(met['PK2'])
    if 'ID_Fund' not in met.columns: met['ID_Fund'] = np.nan
    met['_k'] = key(met)
    cols = ['_k'] + [c for c in met.columns
                     if c in ('JPM_Yield', 'JPM_Duration', 'RA_Yield', 'RA_Duration',
                              'BBG_Yield', 'BBG_Duration', 'BBG_XCCY_Yield')]
    df = df.merge(met[cols].drop_duplicates('_k'), on='_k', how='left')

# Fuentes cuyo Yield sale de una XIRR y por lo tanto llega en DECIMAL
# (0.0817 = 8.17%), a diferencia de JPM/RA/BBG que ya llegan en PORCENTAJE
# (7.86 = 7.86%). Se escalan ×100 al fusionar para que todo el Consolidado
# quede en una sola convención (porcentaje), que es la que usan 06_BREAKEVEN
# y 07_DROPS al leer la columna 'Yield'.
FUENTES_EN_DECIMAL = {'CSHF', 'JSONL', 'EXCEPCIONES'}

def sumar(nombre, path, hoja, col_y, col_d):
    global df
    d = leer(path, hoja)
    if d.empty: return
    d.columns = [str(c).strip() for c in d.columns]
    if 'PK2' not in d.columns or col_y not in d.columns: return
    d['PK2'] = limpiar_txt(d['PK2'])
    if 'ID_Fund' not in d.columns: d['ID_Fund'] = np.nan
    d['_k'] = key(d)
    y = pd.to_numeric(d[col_y], errors='coerce')
    if nombre in FUENTES_EN_DECIMAL:
        y = y * 100   # decimal → porcentaje, misma escala que JPM/RA/BBG
    out = pd.DataFrame({'_k': d['_k'],
                        f'{nombre}_Yield': y,
                        f'{nombre}_Duration': pd.to_numeric(d.get(col_d), errors='coerce')})
    n0 = len(df)
    df = df.merge(out.drop_duplicates('_k'), on='_k', how='left')
    assert len(df) == n0, f"merge con {nombre} multiplicó filas"

sumar('CSHF',        OUT_CSHF,        'resueltos', 'CSHF_Yield_efec', 'CSHF_ModDur')
sumar('JSONL',       OUT_JSONL,       'metricas',  'Yield_efectiva',  'ModDur')
sumar('EXCEPCIONES', OUT_EXCEPCIONES, 'metricas',  'Yield_efectiva',  'ModDur')

for src in CASCADA:
    for suf in ('_Yield', '_Duration'):
        if f'{src}{suf}' not in df.columns:
            df[f'{src}{suf}'] = np.nan

# =============================================================================
# [3] CASCADA
# =============================================================================
print(f"\n[3] Aplicando cascada...")
COL_Y = {s: f'{s}_Yield' for s in CASCADA}
COL_D = {s: f'{s}_Duration' for s in CASCADA}

def elegir(row):
    for s in CASCADA:
        y, d = row.get(COL_Y[s]), row.get(COL_D[s])
        if pd.notna(y) and pd.notna(d):      # TUPLA COMPLETA
            return y, d, s
    return np.nan, np.nan, ''

df[['Yield', 'Duration', 'Fuente']] = df.apply(
    lambda r: pd.Series(elegir(r), index=['Yield', 'Duration', 'Fuente']), axis=1)

# =============================================================================
# [4] DEFAULTEADOS — pisan cualquier valor
# =============================================================================
print(f"\n[4] Aplicando DEF / PROPDEF...")
df['Estado_DEF'] = ''
if os.path.exists(DEFAULTEADOS_PATH):
    dd = pd.read_excel(DEFAULTEADOS_PATH, engine='openpyxl')
    dd.columns = [str(c).strip() for c in dd.columns]
    cdef = next((c for c in dd.columns if c.upper() == 'DEF'), None)
    dd['PK2'] = limpiar_txt(dd['PK2'])
    dd = dd[dd[cdef].astype(str).str.strip().str.upper().isin(MARCAS_DEF)]
    mapa = dict(zip(dd['PK2'], dd[cdef].astype(str).str.strip().str.upper()))
    m = df['PK2'].map(mapa)
    mask = m.notna()
    df.loc[mask, ['Yield', 'Duration', 'Fuente']] = [YIELD_DEF, DURATION_DEF, 'REGLA_DEF']
    df.loc[mask, 'Estado_DEF'] = m[mask]
    print(f"   Posiciones DEF/PROPDEF: {int(mask.sum())} | PK2: {df.loc[mask,'PK2'].nunique()}")
else:
    print(f"   [WARN] falta {os.path.basename(DEFAULTEADOS_PATH)}")

# =============================================================================
# [5] MARCAS
# =============================================================================
df['Metric_Currency']    = df['Risk_Currency']
df['Es_Indexado']        = df['Risk_Currency'].str.upper().isin(INDEXADAS)
df['Requiere_Breakeven'] = df['Es_Indexado'] & (df['Estado_DEF'] == '') & df['Yield'].notna()
df['Requiere_Drop']      = (limpiar_txt(df['Hedge_Currency']) != '') & df['Yield'].notna()
df['Resuelto']           = df['Yield'].notna() & df['Duration'].notna()
df['es_bono']            = pd.to_numeric(df['Issue_Type_Code'], errors='coerce') == ISSUE_TYPE_BONO
df['es_factura']         = df['Name_Instrumento'].str.upper().str.startswith(PREFIJOS_EXCLUIDOS)
df['Estado'] = np.where(df['Estado_DEF'] != '', 'DEFAULT',
                 np.where(df['Resuelto'], 'RESUELTO',
                   np.where(df['es_factura'], 'FACTURA', 'FALTANTE')))

rel = df[~df['es_factura']]        # universo relevante
falt = df[df['Estado'] == 'FALTANTE']
print(f"\n   Resueltos : {int(df['Resuelto'].sum()):,}/{len(rel):,} (universo relevante)")
print(f"   FALTANTES : {len(falt):,}  | PK2 únicos: {falt['PK2'].nunique():,}")
print(f"   Pendientes breakeven: {int(df['Requiere_Breakeven'].sum()):,} | drop: {int(df['Requiere_Drop'].sum()):,}")

# =============================================================================
# [6] HOJAS
# =============================================================================
def pct(a, b): return (a / b) if b else np.nan

COLS_ID = ['ID_Fund', 'FundShortName', 'PK2', 'ISIN', 'Name_Instrumento', 'CompanyName',
           'Risk_Country', 'Risk_Currency', 'Hedge_Currency', 'Issue_Type_Code']

detalle = df[[c for c in COLS_ID if c in df.columns] +
             ['Yield', 'Duration', 'Fuente', 'Metric_Currency', 'Es_Indexado',
              'Requiere_Breakeven', 'Requiere_Drop', 'Estado_DEF', 'Estado'] +
             [f'{s}_Yield' for s in CASCADA] + [f'{s}_Duration' for s in CASCADA] +
             [c for c in ('BBG_XCCY_Yield', 'TotalMVal') if c in df.columns]].copy()
detalle = detalle.sort_values(['FundShortName', 'Estado', 'Name_Instrumento'])

faltantes = df[df['Estado'] == 'FALTANTE'][
    [c for c in COLS_ID if c in df.columns] +
    [c for c in ('TotalMVal',) if c in df.columns]].copy()
faltantes['Tiene_ISIN'] = faltantes['ISIN'] != ''
faltantes['Es_Bono']    = df.loc[faltantes.index, 'es_bono'].values
faltantes = faltantes.sort_values(['FundShortName', 'Name_Instrumento'])

defs   = df[df['Estado_DEF'] != ''][[c for c in COLS_ID if c in df.columns] +
                                    ['Estado_DEF', 'Yield', 'Duration']].copy()
facts  = df[df['es_factura']][[c for c in COLS_ID if c in df.columns] +
                              [c for c in ('TotalMVal',) if c in df.columns]].copy()

# --- Resumen por fondo ---
res = []
for fid, s in rel.groupby('FundShortName'):
    mv_t = pd.to_numeric(s['TotalMVal'], errors='coerce').abs().sum() if 'TotalMVal' in s else np.nan
    mv_o = (pd.to_numeric(s.loc[s['Resuelto'], 'TotalMVal'], errors='coerce').abs().sum()
            if 'TotalMVal' in s else np.nan)
    b = s[s['es_bono']]
    res.append({
        'Fondo': fid, 'ID': int(s['ID_Fund'].dropna().iloc[0]) if s['ID_Fund'].notna().any() else '',
        'Posiciones': len(s), 'Con métricas': int(s['Resuelto'].sum()),
        '% Cobertura': pct(s['Resuelto'].sum(), len(s)),
        'DEF': int((s['Estado'] == 'DEFAULT').sum()),
        'FALTANTES': int((s['Estado'] == 'FALTANTE').sum()),
        '% del AUM cubierto': pct(mv_o, mv_t),
        'Bonos': len(b), '% Bonos cubierto': pct(b['Resuelto'].sum(), len(b)),
    })
res = pd.DataFrame(res).sort_values('% Cobertura', ascending=False).reset_index(drop=True)
tot = {'Fondo': 'TOTAL', 'ID': '', 'Posiciones': len(rel),
       'Con métricas': int(rel['Resuelto'].sum()),
       '% Cobertura': pct(rel['Resuelto'].sum(), len(rel)),
       'DEF': int((rel['Estado'] == 'DEFAULT').sum()),
       'FALTANTES': int((rel['Estado'] == 'FALTANTE').sum()),
       '% del AUM cubierto': pct(
           pd.to_numeric(rel.loc[rel['Resuelto'], 'TotalMVal'], errors='coerce').abs().sum(),
           pd.to_numeric(rel['TotalMVal'], errors='coerce').abs().sum())
           if 'TotalMVal' in rel else np.nan,
       'Bonos': int(rel['es_bono'].sum()),
       '% Bonos cubierto': pct(rel.loc[rel['es_bono'], 'Resuelto'].sum(), rel['es_bono'].sum())}

# --- Panorama ---
uni_rel = rel.drop_duplicates('PK2')
pan = [
    ('UNIVERSO', '', ''),
    ('Posiciones totales (fondo × PK2)', len(df), 'incluye facturas'),
    ('   menos: facturas', -int(df['es_factura'].sum()), 'no son bonos; ningún proveedor las cubre'),
    ('Posiciones del universo relevante', len(rel), 'base de "Resumen por fondo"'),
    ('Papeles ÚNICOS relevantes (PK2)', len(uni_rel),
     f'{len(rel)-len(uni_rel):,} posiciones son el mismo papel en otro fondo'),
    ('', '', ''),
    ('RESULTADO — POSICIONES', '', ''),
    ('Resueltas con métricas', int(rel['Resuelto'].sum()), f"{pct(rel['Resuelto'].sum(), len(rel)):.1%}"),
    ('   de ellas, por regla DEF/PROPDEF', int((rel['Estado'] == 'DEFAULT').sum()), 'Yield=0 / Dur=0.5'),
    ('FALTANTES (sin resolver)', int((rel['Estado'] == 'FALTANTE').sum()),
     f"{pct((rel['Estado']=='FALTANTE').sum(), len(rel)):.1%} → ver hoja FALTANTES"),
    ('', '', ''),
    ('RESULTADO — PAPELES ÚNICOS', '', ''),
    ('Papeles únicos resueltos', int(uni_rel['Resuelto'].sum()),
     f"{pct(uni_rel['Resuelto'].sum(), len(uni_rel)):.1%} de {len(uni_rel):,}"),
    ('Papeles únicos FALTANTES', int((uni_rel['Estado'] == 'FALTANTE').sum()),
     'este es el trabajo real pendiente'),
    ('', '', ''),
    ('FUENTE ELEGIDA (cascada)', '', ' > '.join(CASCADA)),
]
for k, v in rel.loc[rel['Resuelto'], 'Fuente'].value_counts().items():
    pan.append((f'   {k}', int(v), f"{pct(v, rel['Resuelto'].sum()):.1%}"))
pan += [
    ('', '', ''),
    ('CONVERSIONES PENDIENTES', '', ''),
    ('Requieren BREAKEVEN (etapa 06)', int(rel['Requiere_Breakeven'].sum()), 'yield real en moneda indexada'),
    ('Requieren DROP (etapa 07)', int(rel['Requiere_Drop'].sum()), 'papeles con Hedge_Currency'),
]
panorama = pd.DataFrame(pan, columns=['Concepto', 'Valor', 'Comentario'])

# =============================================================================
# [7] ESCRIBIR + FORMATO
# =============================================================================
print(f"\n[7] Guardando {os.path.basename(OUT_CONSOLIDADO)}...")
with pd.ExcelWriter(OUT_CONSOLIDADO, engine='openpyxl') as w:
    panorama.to_excel(w, sheet_name='Panorama', index=False, startrow=3)
    pd.concat([res, pd.DataFrame([tot])], ignore_index=True).to_excel(
        w, sheet_name='Resumen por fondo', index=False, startrow=3)
    detalle.to_excel(w,   sheet_name='Detalle',      index=False, startrow=2)
    faltantes.to_excel(w, sheet_name='FALTANTES',    index=False, startrow=2)
    if len(defs):  defs.to_excel(w,  sheet_name='Defaulteados', index=False, startrow=2)
    if len(facts): facts.to_excel(w, sheet_name='Facturas',     index=False, startrow=2)

wb = load_workbook(OUT_CONSOLIDADO)
TIT = {
    'Panorama': ('Cobertura de métricas — Yield y Duration',
                 f'Cierre {FECHA[6:8]}/{FECHA[4:6]}/{FECHA[:4]}. Un papel está resuelto cuando se obtiene '
                 f'Yield Y Duration de alguna fuente. Cascada: {" > ".join(CASCADA)}.'),
    'Resumen por fondo': ('Resumen por fondo',
                 'Medido en POSICIONES (fondo × PK2), excluidas las facturas.'),
    'Detalle': ('Detalle por posición',
                 'Yield y Duration finales y de qué fuente salieron. Se muestran también los valores de cada fuente.'),
    'FALTANTES': ('Papeles sin resolver — mapeo real',
                 'Universo relevante menos lo resuelto por METRICAS / CSHF / JSONL / EXCEPCIONES y menos los DEF. '
                 'Esto es lo que queda por trabajar.'),
    'Defaulteados': ('Defaulteados (DEF / PROPDEF)',
                 f'Forzados por regla a Yield={YIELD_DEF} y Duration={DURATION_DEF} (liquidación en un semestre).'),
    'Facturas': ('Facturas excluidas del universo',
                 'No son bonos y ningún proveedor de métricas las cubre.'),
}
for sh, (t, sub) in TIT.items():
    if sh not in wb.sheetnames: continue
    ws = wb[sh]; nc = ws.max_column
    hdr = 4 if sh in ('Panorama', 'Resumen por fondo') else 3
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(nc, 3))
    c = ws.cell(1, 1, t); c.font = Font(name=FONT_NAME, size=12, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 20
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(nc, 3))
    ws.cell(2, 1, sub).font = Font(name=FONT_NAME, size=9, italic=True, color='595959')
    for j in range(1, nc + 1):
        h = ws.cell(hdr, j)
        h.fill = PatternFill('solid', fgColor=NAVY)
        h.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
        h.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        h.border = BORDER
    ws.row_dimensions[hdr].height = 28
    for i in range(hdr + 1, ws.max_row + 1):
        for j in range(1, nc + 1):
            cc = ws.cell(i, j)
            cc.font = Font(name=FONT_NAME, size=9); cc.border = BORDER
            v = cc.value
            if isinstance(v, str) and v in ('RESUELTO', 'FALTANTE', 'DEFAULT', 'FACTURA'):
                col = {'RESUELTO': (GREEN_F, GREEN), 'FALTANTE': (RED_F, RED),
                       'DEFAULT': (AMBER_F, '9C5700'), 'FACTURA': (GREY_L, '595959')}[v]
                cc.fill = PatternFill('solid', fgColor=col[0])
                cc.font = Font(name=FONT_NAME, size=9, bold=(v == 'FALTANTE'), color=col[1])
            if isinstance(v, str) and v.isupper() and j == 1 and sh == 'Panorama' and len(v) > 6:
                for jj in range(1, nc + 1):
                    ws.cell(i, jj).fill = PatternFill('solid', fgColor='E8EEF4')
                cc.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)
    if sh == 'Resumen por fondo':
        cols = {c: i + 1 for i, c in enumerate(res.columns)}
        for i in range(hdr + 1, ws.max_row + 1):
            for cn in ('% Cobertura', '% del AUM cubierto', '% Bonos cubierto'):
                ws.cell(i, cols[cn]).number_format = '0.0%'
            v = ws.cell(i, cols['% Cobertura']).value
            if isinstance(v, (int, float)) and i < ws.max_row:
                f_, t_ = ((GREEN_F, GREEN) if v >= .8 else
                          (AMBER_F, '9C5700') if v >= .5 else (RED_F, RED))
                cc = ws.cell(i, cols['% Cobertura'])
                cc.fill = PatternFill('solid', fgColor=f_)
                cc.font = Font(name=FONT_NAME, size=9, bold=True, color=t_)
        for j in range(1, nc + 1):
            cc = ws.cell(ws.max_row, j)
            cc.fill = PatternFill('solid', fgColor=NAVY)
            cc.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
        ws.cell(ws.max_row, cols['% Cobertura']).number_format = '0.0%'
    if sh == 'Detalle':
        dc = {c: i + 1 for i, c in enumerate(detalle.columns)}
        for i in range(hdr + 1, ws.max_row + 1):
            for cn in ('Yield', 'Duration'):
                cc = ws.cell(i, dc[cn]); cc.number_format = '0.0000'
                cc.font = Font(name=FONT_NAME, size=9, bold=True)
            for cn in detalle.columns:
                if cn.endswith(('_Yield', '_Duration')):
                    ws.cell(i, dc[cn]).number_format = '0.0000'
    for j in range(1, nc + 1):
        m = max((len(str(ws.cell(i, j).value)) for i in range(hdr, ws.max_row + 1)
                 if ws.cell(i, j).value is not None), default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(m + 3, 10), 42)
    ws.freeze_panes = ws.cell(hdr + 1, 1).coordinate
    ws.sheet_view.showGridLines = False
    if sh not in ('Panorama', 'Resumen por fondo'):
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(nc)}{ws.max_row}"

wb.save(OUT_CONSOLIDADO)

print(f"\n{'='*70}")
print(f"  Resueltos : {int(rel['Resuelto'].sum()):,}/{len(rel):,}")
print(f"  FALTANTES : {int((rel['Estado']=='FALTANTE').sum()):,}")
print(f"  Output    : {OUT_CONSOLIDADO}")
print(f"{'='*70}")